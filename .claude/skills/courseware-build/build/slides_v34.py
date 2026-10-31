"""v3.4 slides: the founder story. Topic 3 (AI marketing teams), the 5 principles of AI security,
the RAG/evaluation concepts the WA tests, and the bonus Paperclip one-person-company topic.
Topic 2 lives in slides_v34_t2.py. Every slide is written for this course; no repeated row templates."""
import os
from ccdeck import *
import ccdeck as K
from slides_2026 import tag

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.normpath(os.path.join(HERE, '..', '..', '..', '..', 'courseware', 'assets'))
def asset(*p): return os.path.join(ASSETS, *p)
PC = 'https://github.com/paperclipai/paperclip'
HK = 'https://hermes-agent.nousresearch.com/docs/user-guide/features/kanban'
HP = 'https://hermes-agent.nousresearch.com/docs/user-guide/profiles'
HF = 'https://hyperframes.heygen.com/'
BM = 'https://hermes-agent.nousresearch.com/docs/user-guide/bot-mode'

def link(s, x, y, w, label, url, size=10):
    t = text(s, x, y, w, .26, label, size, TEAL, autofit=False); t.text_frame.paragraphs[0].runs[0].hyperlink.address = url

def pills(D, kicker, title, sub, rows_, tipt):
    """The Claude Code course's '5 Principles of AI Security' layout: accent bar · name · dark pill · one-line rule."""
    s = D.new(); header(D, s, kicker, title, sub)
    top, rh, gap = 2.3, .7, .1
    for i, (name, pill, rule) in enumerate(rows_):
        y = top + i * (rh + gap)
        card(s, .6, y, 12.1, rh)
        box(s, .82, y + .2, .07, rh - .4, fill=ACC)
        text(s, 1.1, y, 2.9, rh, name, 15.5, ACC_D, bold=True, font=SERIF, anchor='m', autofit=False)
        box(s, 4.0, y + .2, 2.55, rh - .4, fill=CHAR, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(s, 4.0, y + .2, 2.55, rh - .4, pill.upper(), 9.5, WHITE, bold=True, align='c', anchor='m', autofit=False)
        text(s, 6.85, y + .04, 5.7, rh - .08, rule, 12.5, MUTED, anchor='m', minpt=9.5)
    tip(D, s, tipt, y=top + 5 * (rh + gap) + .02)
    return s

def big_number(D, kicker, number, line, small, dark=True):
    s = D.new(dark=dark)
    circle(s, 9.9, -1.4, 4.2, ACC)
    text(s, .9, 1.7, 9.0, .4, kicker.upper(), 14, ACC, bold=True, autofit=False)
    text(s, .85, 2.15, 9.0, 1.9, number, 96, WHITE if dark else INK, bold=True, font=SERIF, autofit=False)
    text(s, .9, 4.15, 10.5, 1.0, line, 26, WHITE if dark else INK, bold=True, font=SERIF, minpt=16)
    text(s, .9, 5.25, 10.5, 1.2, small, 15, SOFT if dark else MUTED, minpt=11)
    footer(D, s); return s

def before_after(D, kicker, title, sub, left, right, tipt=None):
    """left/right = (label, heading, [lines]) — left muted 'before', right dark 'after'."""
    s = D.new(); header(D, s, kicker, title, sub)
    for j, (lab, hd, lines) in enumerate((left, right)):
        x = .6 + j * 6.15; dark = j == 1
        card(s, x, 2.4, 5.95, 3.95, fill=CHAR if dark else WHITE)
        text(s, x + .35, 2.6, 5.3, .3, lab.upper(), 11, ACC, bold=True, autofit=False)
        text(s, x + .35, 2.95, 5.3, .6, hd, 21, WHITE if dark else INK, bold=True, font=SERIF, minpt=13)
        text(s, x + .35, 3.7, 5.3, 2.5, "\n".join(lines), 14.5, SOFT if dark else MUTED, lh=1.3, minpt=10)
    line(s, 6.45, 4.35, 6.85, 4.35, ACC, 3, arrow=True)
    if tipt: tip(D, s, tipt, y=6.45)
    return s

# ====================================================================== Topic 3
def topic3_open(D, C):
    section(D, 'Day 2 · Topic 3', C.TOPICS[2]['title'], C.TOPICS[2]['items'],
            sub='Recap: a CEO assistant, then your first AI employee. Today: AI teams that run on a schedule.')
    big_number(D, 'Topic 3 · Two weeks to launch', '5 + 5',
               'Five bots in a group chat. Five agents on a board.',
               'Yesterday Mia worked alone. Today you hire a team: five named bots write the launch posts together in one group chat, '
               'one prompt builds a video team on a Kanban board, and a cron sets both to work every week. You approve everything.')

def video_tools(D):
    s = D.new(); header(D, s, 'Topic 3 · Video skills and tools', 'Three Ways Your Agents Make Video with Code',
                        'No camera, no editor: the agent writes code, the tool renders an MP4. Pick the one that suits the job.')
    tools = [('HyperFrames', 'by HeyGen · HTML → MP4', 'Write each scene as HTML/CSS with GSAP timing; lint, snapshot, preview, render. Narrated with Hermes’s free Edge voice.',
              'Kopi Kaki uses it for the launch video.', 'Ask Mia: “Install the official\nhyperframes skill and check\nwhat it needs.”', HF),
             ('Remotion', 'React → MP4', 'Each frame is a React component; strong for data-driven and templated videos at scale.',
              'Good for 50 personalised promo cuts.', 'Capabilities → Skills →\nsearch “remotion”', 'https://www.remotion.dev/'),
             ('Manim', 'Python → MP4', 'The maths-animation engine behind 3Blue1Brown: precise diagrams, graphs and equations that move.',
              'Good for explaining how cold brew is made.', 'Built in: the manim-video\nskill (Mia installs Manim\nif you ask)', 'https://www.manim.community/')]
    for i, (name, tag_, what, use, cmd, url) in enumerate(tools):
        x = .6 + i * 4.1; dark = False
        card(s, x, 2.35, 3.9, 4.35, fill=CHAR if dark else WHITE)
        text(s, x + .3, 2.5, 3.3, .5, name, 22, WHITE if dark else INK, bold=True, font=SERIF, autofit=False)
        text(s, x + .3, 3.02, 3.3, .3, tag_, 11.5, ACC, bold=True, autofit=False)
        text(s, x + .3, 3.4, 3.3, 1.25, what, 12, SOFT if dark else MUTED, minpt=9.5)
        text(s, x + .3, 4.62, 3.3, .45, use, 12, WHITE if dark else INK, bold=True, minpt=9.5)
        card(s, x + .25, 5.12, 3.4, .95, fill=K.CODEBG, shadow=False)
        text(s, x + .38, 5.18, 3.2, .85, cmd, 10, 'F1EDE4', font=K.MONO, minpt=8, cw=.62)
        link(s, x + .3, 6.32, 3.4, url.replace('https://', '') + ' · click', url)
    tag(D, 'TOOLS', HF + ' ; https://www.remotion.dev/ ; https://www.manim.community/')

def bots_and_groups(D):
    """Activity 4 concepts: hire named bots, then put them in one group chat (Hermes Desktop Bot Mode)."""
    image(D, 'Topic 3 · Bots', 'Hire a Bot in 30 Seconds: Name, Title, Role', asset('screenshots', 'desktop', 'bot-edit-profile.png'),
          side=[('users', 'Bots tab → New Agent', 'Three fields: Name, Title, Description. The bot introduces itself.'),
                ('file', 'Advanced', 'Paste its role file as Custom SOUL.md; pick its model, skills and tools.'),
                ('sparkles', 'Give it a face', 'Right-click → Edit Profile: Bot, Generate, Upload or a Pet.')],
          caption='A bot is a Hermes profile with a name you can @mention: its own role, model, memory and skills, and one permanent Bot Chat.')
    tag(D, 'ARCH', BM)
    image(D, 'Topic 3 · Group chats', 'Put Your Bots in One Room and Let Them Hand Off', asset('screenshots', 'desktop', 'bot-group-chat.png'),
          side=[('users', 'Manage groups', 'Right-click a bot → Manage groups → create one. 2–6 bots per room.'),
                ('chat', '@name hands off', 'Bots pull teammates in with @name; @user asks you to decide.'),
                ('shield', 'Built-in brakes', 'Up to 3 rounds and 10 messages per send. Stop, Rename, Disband.')],
          caption='Trainer’s test room: Eve and Dalong answer the same message and tag each other. Your room has five marketers.')
    tag(D, 'ARCH', BM)

def social_team(D):
    s = D.new(); header(D, s, 'Topic 3 · Activity 4', 'Meet the Kopi Social Team: Five Names, Five Jobs',
                        'Give each bot a name you will remember and a role it cannot misread. They share one group chat; you have the final word.')
    team = [('search', 'Rafi', 'Market Researcher', 'Angles and facts, each with its source file. Goes first.'),
            ('linkedin', 'Lin', 'LinkedIn Writer', 'The founder story for investors and catering buyers.'),
            ('users', 'Farah', 'Facebook Writer', 'Neighbours, pop-up dates and the 1 Dec opening.'),
            ('image', 'Ivy', 'Instagram Writer', 'The scroll-stopping hook, caption and hashtags.'),
            ('sparkles', 'Pixel', 'Image Creator', 'One on-brand image per post and the review gallery.'),
            ('user', 'You', 'Founder · @user', 'APPROVE, EDIT or REJECT every post. Nothing goes out without you.')]
    for i, (ic, n, ttl, d) in enumerate(team):
        r, c = divmod(i, 3); x = .6 + c * 4.1; y = 2.35 + r * 2.05; you = n == 'You'
        card(s, x, y, 3.9, 1.85, fill=ACC_D if you else (CHAR if r == 0 and c == 0 else WHITE))
        dark = you or (r == 0 and c == 0)
        icon(s, ic, x + .25, y + .28, .6, fill=WHITE if you else ACC, glyph=ACC_D if you else WHITE)
        text(s, x + 1.05, y + .2, 2.7, .45, n, 21, WHITE if dark else INK, bold=True, font=SERIF, autofit=False)
        text(s, x + 1.05, y + .66, 2.7, .3, ttl.upper(), 10.5, LIGHT if you else ACC, bold=True, autofit=False)
        text(s, x + .25, y + 1.05, 3.45, .75, d, 12, SOFT if dark else MUTED, minpt=9)
    tag(D, 'CONCEPT', 'Course design; role files in labs/lab04-social-media-marketing-team/data/souls/')

def kanban_board(D):
    s = D.new(); header(D, s, 'Topic 3 · Kanban', 'Parents First: A Card Runs Only When Its Inputs Exist',
                        'The board is the video team’s shared to-do list. One prompt to Mia creates the team and these seven cards; your two HUMAN cards hold everything after them.')
    nodes = [('Research', 'video-researcher'), ('Script', 'scriptwriter'), ('HUMAN', 'approve script'), ('Voice-over', 'voice-over'),
             ('Render', 'video-producer'), ('Review', 'video-reviewer'), ('HUMAN', 'approve exact video')]
    w, gap, y = 1.55, .2, 3.3
    for i, (lab, who) in enumerate(nodes):
        x = .6 + i * (w + gap); human = lab == 'HUMAN'
        card(s, x, y, w, 1.35, fill=ACC_D if human else WHITE)
        text(s, x + .1, y + .15, w - .2, .45, lab, 15, WHITE if human else INK, bold=True, font=SERIF, align='c', autofit=False)
        text(s, x + .1, y + .65, w - .2, .6, who, 10.5, LIGHT if human else MUTED, align='c', minpt=8)
        if i < len(nodes) - 1: line(s, x + w + .01, y + .68, x + w + gap - .01, y + .68, ACC, 1.75, arrow=True)
    text(s, .6, 5.05, 12.1, .9, 'Bots in a group chat (Activity 4) talk their way to a result. A board (Activity 5) makes the order explicit: '
         'who goes next, what each card needs, and where the work stops for you. Use a board when the steps are known and the stakes are high.', 14, MUTED, minpt=11)
    tip(D, s, 'Kanban in the left sidebar shows the board live. Open a card for its run history and output; PP Task 3 screenshot A comes from here.', y=6.2)
    tag(D, 'ARCH', HK)
    image(D, 'Topic 3 · Kanban', 'Watch the Board Work', asset('screenshots', 'hermes-kanban.png'),
          caption='Reference Kanban view: cards move from waiting to running to done as their parents finish. Your board shows the kopi-video cards and your HUMAN gates.')

def activity4_team(D):
    s = D.new(dark=True); header(D, s, 'Topic 3 · Activity 4', 'One Brief, Three Platforms, Three Voices',
                                 'The same launch fact, written by the specialist for each audience.')
    posts = [('LinkedIn', 'linkedin', 'Founder story', '“I quit my bank job to bring back the kopitiam I grew up with. On 1 Dec our Tiong Bahru flagship opens…”'),
             ('Facebook', 'users', 'The neighbourhood', '“Tiong Bahru kakis! Catch our last pop-up this Saturday, then come say hi on 1 Dec. Gula Melaka Cold Brew, SGD 6.80.”'),
             ('Instagram', 'image', 'The scroll-stopper', '“Slow-brewed. Shiok cold. 12 hours, real gula melaka, coconut cream. #ShiokColdBrew”')]
    for i, (pl, ic, aud, q) in enumerate(posts):
        x = .6 + i * 4.1
        card(s, x, 2.45, 3.9, 3.7, fill=K.PANEL, shadow=False); icon(s, ic, x + .3, 2.7, .56)
        text(s, x + 1.05, 2.68, 2.6, .4, pl, 19, WHITE, bold=True, font=SERIF, autofit=False)
        text(s, x + 1.05, 3.08, 2.6, .3, aud, 11.5, ACC, bold=True, autofit=False)
        text(s, x + .3, 3.65, 3.3, 2.3, q, 14, SOFT, font=SERIF, minpt=10, lh=1.2)
    tip(D, s, 'Sample copy for illustration. Your team writes its own from the brief, and you approve every post.', y=6.35)

def video_team(D):
    s = D.new(); header(D, s, 'Topic 3 · Activity 5', 'Your Production Line: Five Agents, Two Human Gates',
                        'Research → script → YOU approve → voice-over and render → independent review → YOU approve the exact file.')
    steps = [('search', 'Researcher', 'Claim | source | quote', False), ('pen', 'Scriptwriter', '4 scenes · 20 s', False), ('user', 'You', 'Approve the script', True),
             ('chat', 'Voice-over', 'Free Edge voice', False), ('film', 'Producer', 'HyperFrames → MP4', False), ('eyecheck', 'Reviewer', 'V1–V9 + link', False), ('user', 'You', 'APPROVE <sha256>', True)]
    w = 1.58; gap = .14
    for i, (ic, t, d, human) in enumerate(steps):
        x = .6 + i * (w + gap)
        card(s, x, 2.6, w, 2.7, fill=ACC_D if human else WHITE)
        icon(s, ic, x + .5, 2.85, .56, fill=WHITE if human else ACC, glyph=ACC_D if human else WHITE)
        text(s, x + .1, 3.6, w - .2, .45, t, 14.5, WHITE if human else INK, bold=True, font=SERIF, align='c', minpt=10)
        text(s, x + .1, 4.1, w - .2, 1.0, d, 11.5, LIGHT if human else MUTED, align='c', minpt=9)
        if i < len(steps) - 1: line(s, x + w + .01, 3.95, x + w + gap - .01, 3.95, ACC, 1.5, arrow=True)
    text(s, .6, 5.6, 12.1, .7, 'The reviewer sends you a link (http://localhost:8765/final.mp4) and the file’s SHA-256. Your approval belongs to that exact file: change one frame and it needs a new review.', 14, MUTED, minpt=11)
    tag(D, 'ARCH', HK + ' ; ' + HF)

# ---------------------------------------------------------------- grounding / RAG (WA K2) — before Activity 5
def grounding(D):
    before_after(D, 'Topic 3 · Grounding', 'Grounded or Guessing? RAG versus a Standard LLM',
                 'Your researcher bot is a small RAG system: it may only say what a source file supports.',
                 ('Standard LLM', 'Answers from what it learned', ['Fluent, fast, cheap to set up.', 'Does not know your price, dates or latest policy.', 'Can state a wrong fact confidently.']),
                 ('Retrieval-augmented (RAG)', 'Answers from passages it retrieved', ['Quotes the approved source and its version.', 'Says INSUFFICIENT EVIDENCE when nothing supports a claim.', 'Needs owners, indexing, access control and refresh.']),
                 tipt='In Activity 5 the researcher writes claim | source file | exact quote. That table is your RAG audit trail.')
    tag(D, 'CONCEPT', 'Lewis et al. 2020, Retrieval-Augmented Generation; course design')
    cards(D, 'Topic 3 · Grounding', 'Four Small Tests Reveal Different Failure Modes', [
        ('Supported', '“What does the cold brew cost?” → the catalogue line: SGD 6.80.', 'check'),
        ('Partial', '“What is the SHIOK20 offer?” → the discount AND its terms, or it is incomplete.', 'file'),
        ('Unsupported', '“Is it healthier than soft drinks?” → INSUFFICIENT EVIDENCE. No source, no claim.', 'shield'),
        ('Conflicting', 'Two files disagree on the launch date → stop and ask the source owner.', 'alert')],
        cols=4, sub='Run all four against your researcher before you trust it with a public video.')
    tag(D, 'EVIDENCE', 'Course test design; Kopi Kaki source files')
    s = D.new(); header(D, s, 'Topic 3 · Grounding', 'Refresh Sources Without Reusing Obsolete Advice',
                        'A RAG system is only as current as its last refresh. Maintenance is part of feasibility, not an afterthought.')
    for i, (n, t, d) in enumerate([('1', 'Change', 'The owner approves price list v2 (SGD 6.80 → 7.20).'), ('2', 'Refresh', 'Re-index the changed passages; retire v1 so it cannot be retrieved.'),
                                   ('3', 'Retest', 'Re-run the four tests: supported, partial, unsupported, conflicting.'), ('4', 'Re-review', 'Every post and video that quoted v1 goes back to review.')]):
        x = .6 + i * 3.08
        text(s, x, 2.5, 2.9, 1.2, n, 72, ACC, bold=True, font=SERIF, autofit=False)
        text(s, x, 3.8, 2.9, .5, t, 19, INK, bold=True, font=SERIF)
        text(s, x, 4.35, 2.85, 1.6, d, 13.5, MUTED, minpt=10)
    tip(D, s, 'For a regulated business (bank, insurer, clinic) add a named source owner, access control per document and an audit log of every refresh.', y=6.2)
    tag(D, 'CONCEPT', 'Course design')
    table(D, 'Topic 3 · Grounding', 'A Retrieval Evaluation Must Measure More Than Fluency', ['Measure', 'Question it answers', 'How to check it in Activity 5'],
          [('Relevance', 'Did search return the passage needed?', 'Each claim in research.md points to the right file'),
           ('Support', 'Does the passage actually justify the answer?', 'Read the quoted line yourself for two claims'),
           ('Latency', 'How long from question to reviewed answer?', 'Card start → reviewer PASS, from run history'),
           ('Cost', 'Ingestion + queries + human review', 'Model spend plus your review minutes')], colw=[2.2, 4.6, 5.3], size=13)
    tag(D, 'EVIDENCE', 'Course evaluation design')

# ---------------------------------------------------------------- Activity 6 · content crons
def crons_week(D):
    s = D.new(dark=True); header(D, s, 'Topic 3 · Activity 6', 'A Week on Autopilot, With You Holding the Keys',
                                 'Two scheduled jobs start the team’s work. Neither one approves anything.')
    week = [('MON 9:00', 'Weekly social drafts', 'Three posts, three images and a gallery in a dated folder.', False),
            ('MON 9:05', 'Telegram ping', '“Drafts ready — APPROVE, EDIT or REJECT.”', False),
            ('YOU', 'Approve on your phone', 'Only approved posts are used.', True),
            ('FRI 10:00', 'Weekly video draft', 'A new run on the kopi-video board.', False),
            ('YOU', 'Two gates', 'Approve the script, then the exact video.', True)]
    w, gap = 2.3, .15
    for i, (when, t, d, you) in enumerate(week):
        x = .6 + i * (w + gap)
        card(s, x, 2.45, w, 3.1, fill=ACC_D if you else K.PANEL, shadow=False)
        text(s, x + .2, 2.65, w - .4, .4, when, 15, WHITE if you else ACC, bold=True, autofit=False)
        text(s, x + .2, 3.15, w - .4, .9, t, 17, WHITE, bold=True, font=SERIF, minpt=12)
        text(s, x + .2, 4.15, w - .4, 1.3, d, 12.5, LIGHT if you else SOFT, minpt=9.5)
    text(s, .6, 5.8, 12.1, .5, 'Automate the start of the work, never the approval. Test with Trigger now; pause what you cannot watch.', 14, LIGHT, minpt=11)
    tag(D, 'CONCEPT', 'Course design; ' + 'https://hermes-agent.nousresearch.com/docs/user-guide/features/cron')

# ---------------------------------------------------------------- evaluation (WA K3/K5) — after Activity 5
def evaluation(D):
    cards(D, 'Topic 3 · Evaluate and improve', 'Four Metrics Support an Improvement Decision', [
        ('Cycle time', 'Brief to accepted video, including your reviews and every retry.', 'clock'),
        ('First-pass acceptance', 'Share of outputs you approved without a revision.', 'check'),
        ('Cost per accepted item', 'Model and tool spend plus review minutes, per approved post or video.', 'coins'),
        ('Defect rate', 'Fact, brand and video-quality failures, counted separately.', 'alert')],
        cols=2, sub='Measure the same four numbers before and after any change to a skill, a SOUL.md or the team.')
    tag(D, 'EVIDENCE', 'Course evaluation design')

# ---------------------------------------------------------------- 5 principles of AI security
def security(D):
    pills(D, 'AI security', '5 Principles of AI Security',
          'Your agents read your inbox, spend your money and speak for your brand. These five controls keep what they see and do safe, reviewable and accountable.',
          [('Data Control', 'What AI can see', 'Classify your data (public, internal, confidential) and decide exactly which data the AI may access.'),
           ('Permission Control', 'What AI can do', 'Grant only the rights needed — read, write, edit, delete. Ask before anything destructive.'),
           ('Human in the Loop', 'Who approves', 'A person reviews and approves high-impact actions — sends, posts, payments, deletions.'),
           ('Traceability', 'What happened', 'Log every prompt, tool call and change so actions can be audited, explained and rolled back.'),
           ('Right-Sized Models', 'Not the most powerful', 'Do not use the most powerful model by default — the smallest model that does the job has less autonomy and risk.')],
          'In Hermes: a read-only Gmail connector, a Telegram approved-users list, @user questions in group chats, HUMAN Kanban cards, file fingerprints, a smaller model per bot.')
    tag(D, 'POLICY', 'Same five principles as TGS-2025052468 Agentic AI Applications with Claude Code (slide 75)')

# ====================================================================== Bonus · Paperclip
def bonus(D, C, mark):
    mark('bonus')
    section(D, 'Bonus topic · optional', C.BONUS['title'].replace('Bonus (optional) — ', ''), C.BONUS['items'],
            sub='One founder. An AI company. You sit on the Board.')
    big_number(D, 'Bonus · The one-person company', '1 human', 'Kopi Kaki Co. decides to run as a one-person company.',
               'You are the founder, the Board and the only employee with a pulse. An AI CEO runs the day-to-day, hires a Chief of Staff, '
               'and the Chief of Staff hires the marketing team. Your job: set the mission, approve the hires, approve what goes public.')
    before_after(D, 'Bonus · What is Paperclip?', 'Kanban Organises Tasks. Paperclip Organises People.',
                 'Paperclip (open source) runs a company of AI agents: org chart, job titles, reporting lines, budgets and an approval inbox.',
                 ('Hermes Kanban (Activity 5)', 'You are the manager', ['You create every card and every dependency.', 'You decide who does what.', 'Budgets and hiring are not modelled.']),
                 ('Paperclip company (bonus)', 'You are the Board', ['A CEO agent breaks goals into tasks and hires.', 'Hires wait in your approval inbox.', 'Each agent has a budget and a manager.']),
                 tipt='Paperclip runs on agents you already have: Hermes, Claude Code, Codex or Gemini CLI become its “adapters”.')
    tag(D, 'ARCH', PC)
    image(D, 'Bonus · Install', 'Install Paperclip: One Request to Mia, Then localhost:3100', asset('paperclip', 'dashboard.png'),
          caption='Ask Mia to install and start it, then open http://localhost:3100. Trainer’s demo company shown; you build Kopi Kaki Co. the same way.',
          side=[('gear', 'Prerequisite', 'Node.js 24.11+ on the laptop (the trainer installs it) and Mia from Topic 2.'),
                ('chat', 'Install', 'Ask Mia: “Install Paperclip and keep it running.” She runs the installer for you.'),
                ('building', 'Company', 'Wizard: name it Kopi Kaki Co.; first agent Ava.')])
    image(D, 'Bonus · Set up the company', 'Kopi Kaki Co.: Mission, Budget, Hiring Gate', asset('paperclip', 'company-settings.png'),
          caption='Company settings hold the name, mission and budget. Then Settings → General → Hiring → “Require board approval for new hires” = ON.')
    s = D.new(); header(D, s, 'Bonus · The CEO', 'Write the CEO’s Job Description Before You Give Her Work',
                        'Agents → Ava → Instructions → AGENTS.md. These instructions steer every decision she makes.')
    card(s, .6, 2.35, 7.7, 4.4, fill=K.CODEBG)
    text(s, .85, 2.5, 7.3, 4.15, 'You are Ava, CEO of Kopi Kaki Co., a Singapore kopi startup run as a one-person company:\nthe founder is the Board and the only human.\n\n'
         'Mission: launch the Gula Melaka Cold Brew on 1 Dec 2026.\n\n- You do not do the marketing work. You hire and manage.\n- First hire: a Chief of Staff who runs the marketing team.\n'
         '- Every hire: mandate, reporting line, budget, Board approval.\n- Every public item: independent reviewer, then the Board.\n- Report weekly: progress, spend, risks, decisions.',
         12.5, 'F1EDE4', font=K.MONO, minpt=9, cw=.62, lh=1.12)
    for i, (ic, t, d) in enumerate([('briefcase', 'Role: CEO', 'Paperclip has no Role field: ask Mia to set it'), ('file', 'Source files', 'Upload brand guide, brief, compliance rules'),
                                    ('gavel', 'First task', '“Hire a Chief of Staff” — it lands in your approval inbox')]):
        y = 2.45 + i * 1.45; icon(s, ic, 8.65, y, .52)
        text(s, 9.3, y - .03, 3.4, .4, t, 15, INK, bold=True, font=SERIF); text(s, 9.3, y + .4, 3.4, .9, d, 12, MUTED, minpt=9)
    tag(D, 'PROMPT', 'Bonus lab prompt; ' + PC)
    flow(D, 'Bonus · Delegated hiring', 'The CEO Hires a Chief of Staff; the Chief of Staff Hires the Team', [
        ('user', 'You', 'Task to Ava: hire a Chief of Staff.'), ('briefcase', 'CEO', 'Proposes mandate, budget, reporting line.'),
        ('gavel', 'Board', 'You approve in Pending Approvals.'), ('users', 'Chief of Staff', 'Proposes six marketing hires.'),
        ('gavel', 'Board', 'Approve, revise or reject each.')],
        sub='Managers propose; only the Board approves. That is how a one-person company keeps control.')
    image(D, 'Bonus · Org chart', 'Your One-Person Company on One Page', asset('paperclip', 'org-chart.png'),
          caption='Agents → Org chart view: Board → CEO → Chief of Staff → marketing team. Trainer’s demo company shown; yours shows the Kopi Kaki roles.')
    image(D, 'Bonus · Deliver', 'The Team Produces the Launch Video', asset('paperclip', 'tasks-board.png'),
          caption='Give the Chief of Staff the Activity 5 brief. Sub-tasks move through Todo → In Progress → In Review → Done, with your approvals in between.')
    table(D, 'Bonus · Compare', 'Same Video, Two Ways: Hermes Kanban vs Paperclip', ['', 'Hermes Kanban team (Activity 5)', 'Paperclip company (bonus)'],
          [('Who breaks down the work', 'You create every card and parent', 'The CEO and Chief of Staff create tasks'),
           ('Who decides the team', 'You (one prompt to Mia) create the agents', 'Managers propose hires; you approve'),
           ('Control', 'HUMAN cards on the board', 'Approval inbox, budgets, hiring gate'),
           ('Visibility', 'Board, run history, files', 'Org chart, audit log, costs per agent'),
           ('Best for', 'A known pipeline you want to run exactly', 'An ongoing operation you want to delegate and govern')],
          colw=[2.6, 4.75, 4.75], size=13, tipt='Write comparison.md in the bonus lab: which would you use for Kopi Kaki next month, and why?')
    tag(D, 'EVIDENCE', 'Learner comparison; ' + PC + ' ; ' + HK)

# ---------------------------------------------------------------- multi-agent communication (trainer slides, v3.3 hand edit → native)

def multi_intro(D):
    """General introduction to multiple agents (trainer request): why a team, and what every agent owns."""
    before_after(D, 'Topic 3 · Multiple agents', 'From One Employee to a Team of Agents',
                 'A multi-agent system splits one big goal into specialist jobs that hand work to each other — with a human in charge.',
                 ('One agent (Topic 2)', 'Mia does everything', ['One context window fills up with every task.', 'One set of instructions for every voice and job.', 'Jobs run one after another.']),
                 ('Multiple agents (Topic 3)', 'A team of specialists', ['Each agent sees only what its job needs.', 'Each has its own instructions, skills and tools.', 'Independent jobs run in parallel; a board tracks hand-offs.']),
                 tipt='More agents also means more hand-offs to check. Use a team when the jobs are truly different, not just because you can.')
    tag(D, 'CONCEPT', 'Course design; Hermes profiles and Kanban docs')
    s = D.new(); header(D, s, 'Topic 3 · Multiple agents', 'Every Agent Brings Its Own Five Things',
                        'In Hermes each agent is a bot (a profile). Change any of the five and you get a different employee.')
    five = [('tool', 'Tools', 'What it can do', 'web, terminal, image, voice …', 'Pixel: image tool on'),
            ('puzzle', 'Unique skills', 'How it does the job well', 'its own SKILL.md set', 'video-producer: hyperframes'),
            ('brain', 'Model', 'How it thinks', 'big or small, any provider', 'Rafi: a careful model'),
            ('file', 'Instructions', 'Who it is and its rules', 'SOUL.md + AGENTS.md', 'reviewer: never approves for you'),
            ('layers', 'Context window', 'What it is holding in mind', 'its own memory and session', 'each bot keeps its own chat and memory')]
    for i, (ic, t, q, how, eg) in enumerate(five):
        x = .6 + i * 2.46; dark = i % 2 == 1
        card(s, x, 2.4, 2.3, 4.0, fill=CHAR if dark else WHITE)
        icon(s, ic, x + .85, 2.65, .62, fill=ACC)
        text(s, x + .15, 3.45, 2.0, .45, t, 17, WHITE if dark else INK, bold=True, font=SERIF, align='c', minpt=12)
        text(s, x + .15, 3.95, 2.0, .55, q, 12, ACC, bold=True, align='c', minpt=9)
        text(s, x + .15, 4.55, 2.0, .75, how, 12, SOFT if dark else MUTED, align='c', minpt=9)
        text(s, x + .15, 5.35, 2.0, .9, eg, 11, LIGHT if dark else INK, align='c', minpt=8.5)
    tip(D, s, 'Hermes Desktop: Bots tab → New Agent → Name, Title, Description; Advanced for its SOUL.md, skills, tools and model.', y=6.55)
    tag(D, 'ARCH', HP)

def each_can(D):
    s = D.new(dark=True); header(D, s, 'Topic 3 · How agents talk', 'Each Agent Can…',
                                 'That is why a team beats one big agent: every member is set up differently. In Hermes, each bot is one member.')
    for i, (ic, t, kk) in enumerate([('tool', 'have different tools', 'Pixel has the image tool; Rafi has web search.'),
                                     ('brain', 'access different knowledge', 'Farah knows the pop-up dates; Lin knows the founder story.'),
                                     ('lock', 'operate under different permissions', 'reviewer can read but not render; nobody can post — only you approve.'),
                                     ('users', 'perform distinct roles', 'each SOUL.md says what the role owns and must never do.')]):
        x = .6 + i * 3.08
        card(s, x, 2.5, 2.9, 3.85, fill=K.PANEL, shadow=False)
        icon(s, ic, x + .95, 2.8, 1.0)
        text(s, x + .2, 4.0, 2.5, .8, t, 17, WHITE, bold=True, font=SERIF, align='c', minpt=12)
        text(s, x + .25, 4.9, 2.4, 1.35, 'At Kopi Kaki: ' + kk, 12, SOFT, align='c', minpt=9)
    tag(D, 'CONCEPT', 'Trainer slide (v3.3 hand edit) recreated; Hermes profiles docs')

def handoff(D):
    s = D.new(dark=True); header(D, s, 'Topic 3 · How agents talk', '“My order arrived damaged. I need a replacement.”',
                                 'One customer message, four agents, one outcome. Each hands a structured message to the next.')
    ag = [('user', 'Support Agent', 'Understands the complaint, checks the order, opens a case.'),
          ('box', 'Logistic Agent', 'Checks stock and reserves a replacement item.'),
          ('send', 'Shipping Agent', 'Books the courier and returns a tracking number.'),
          ('loop', 'Customer Agent', 'Confirms the fix to the customer and asks for a rating.')]
    w = 2.75; gap = .37
    for i, (ic, t, d) in enumerate(ag):
        x = .6 + i * (w + gap)
        card(s, x, 2.55, w, 3.2, fill=K.PANEL, shadow=False)
        icon(s, ic, x + .85, 2.8, 1.05)
        text(s, x + .15, 4.0, w - .3, .45, t, 17, WHITE, bold=True, font=SERIF, align='c', minpt=12)
        text(s, x + .2, 4.5, w - .4, 1.15, d, 12.5, SOFT, align='c', minpt=9.5)
        if i < len(ag) - 1: line(s, x + w + .03, 4.15, x + w + gap - .03, 4.15, ACC, 2.25, arrow=True)
    tip(D, s, 'At Kopi Kaki: “My cold brew pre-order arrived spilled.” Same four hand-offs. A refund above SGD 20 waits for you: human in the loop.', y=6.1)
    tag(D, 'CONCEPT', 'Trainer slide (v3.3 hand edit) recreated')

HD = 'https://hermes-agent.nousresearch.com/docs/user-guide/features/delegation'
# Topic 3 concept slides: the trainer's v3.5 multi-agent screenshots redrawn natively (same ideas, Kopi Kaki examples).

def hub_spoke(D):
    s = D.new(dark=True); header(D, s, 'Topic 3 · How agents work together', 'Hub-and-Spoke: One Coordinator, Many Specialists',
                                 'No single agent carries the whole job. The hub hands out the pieces, the spokes do them, and the hub puts the answer together.')
    card(s, .6, 2.4, 4.45, 3.9, fill=K.PANEL, shadow=False)
    text(s, .85, 2.58, 4.0, .28, 'THE IDEA', 11, ACC, bold=True, autofit=False)
    text(s, .85, 2.9, 4.0, 1.0, 'Don’t give every job to one agent. Split the work among specialists who all pull toward one goal.', 16, WHITE, font=SERIF, minpt=12, lh=1.15)
    for i, (ic, t, d) in enumerate([('robot', 'Hub · the coordinator', 'Hands out the pieces, collects the results, answers you.'),
                                    ('users', 'Spokes · the specialists', 'Each does one piece with its own skills and tools.'),
                                    ('target', 'One shared goal', 'Every piece comes back to a single answer.')]):
        y = 4.05 + i * .72
        icon(s, ic, .85, y + .04, .46)
        text(s, 1.45, y, 3.45, .3, t, 13, WHITE, bold=True, font=SERIF, autofit=False)
        text(s, 1.45, y + .3, 3.45, .38, d, 11, SOFT, minpt=9)
    hx, hy, hd = 5.75, 3.55, 1.45
    spokes = [('search', 'Rafi', 'market research'), ('linkedin', 'Lin', 'LinkedIn post'), ('users', 'Farah', 'Facebook post'),
              ('image', 'Ivy', 'Instagram post'), ('camera', 'Pixel', 'post images')]
    for i, (ic, n, job) in enumerate(spokes):
        y = 2.4 + i * .8
        line(s, hx + hd, hy + hd / 2, 9.25, y + .31, SOFT, 1.5)
        card(s, 9.25, y, 3.45, .62, fill=K.PANEL, shadow=False)
        icon(s, ic, 9.35, y + .08, .46, fill=TEAL)
        text(s, 9.92, y + .06, 1.0, .5, n, 14, WHITE, bold=True, font=SERIF, anchor='m', autofit=False)
        text(s, 10.85, y + .06, 1.8, .5, job, 11.5, SOFT, anchor='m', minpt=9)
    icon(s, 'user', hx, hy, hd)
    text(s, hx - .45, hy + hd + .1, hd + .9, .3, 'Coordinator (hub)', 13, WHITE, bold=True, font=SERIF, align='c', autofit=False)
    text(s, hx - .45, hy + hd + .42, hd + .9, .3, 'Activity 4: you', 11, ACC, bold=True, align='c', autofit=False)
    tip(D, s, 'Activity 5 has the same shape: Mia is the hub; the researcher, scriptwriter, voice-over, producer and reviewer are the spokes.', y=6.5, label='SAME SHAPE')
    tag(D, 'CONCEPT', 'Redrawn from trainer-supplied screenshots (v3.5); ' + HK)

def comm_flow(D):
    s = D.new(dark=True); header(D, s, 'Topic 3 · How agents work together', 'Out to the Spokes, Back to the Hub',
                                 'Every request makes the same five-step round trip. Here it is with your Activity 5 video team.')
    steps = [('chat', 'Request in', 'Your request reaches the coordinator.', '“Make a 20-second launch video for the Gula Melaka Cold Brew.”'),
             ('layers', 'Split', 'The coordinator breaks it into tasks.', 'Mia creates seven cards on the kopi-video board.'),
             ('send', 'Fan out', 'Each task goes to a specialist.', 'Researcher, scriptwriter, voice-over, producer, reviewer.'),
             ('file', 'Results back', 'Each specialist returns its output.', 'research.md, script.md, narration, MP4, review notes.'),
             ('check', 'One answer', 'The coordinator answers you once.', 'Mia reports; nothing ships until you approve the exact video.')]
    w, gap = 2.3, .15
    for lab, i0, i1, col in [('OUT TO THE SPOKES  →', 0, 2, ACC), ('←  BACK TO THE HUB', 3, 4, TEAL)]:
        x0 = .6 + i0 * (w + gap); x1 = .6 + i1 * (w + gap) + w
        box(s, x0, 2.38, x1 - x0, .06, fill=col)
        text(s, x0, 2.48, x1 - x0, .3, lab, 11, col, bold=True, align='c', autofit=False)
    for i, (ic, t, g, ex) in enumerate(steps):
        x = .6 + i * (w + gap); cx = x + w / 2
        icon(s, ic, cx - .42, 2.95, .84, fill=ACC if i < 3 else TEAL)
        if i < 4: line(s, cx + .5, 3.37, cx + w + gap - .5, 3.37, SOFT, 1.75, arrow=True)
        text(s, x, 3.9, w, .38, f'0{i + 1} · {t}', 15, WHITE, bold=True, font=SERIF, align='c', autofit=False)
        text(s, x + .05, 4.3, w - .1, .6, g, 12, SOFT, align='c', minpt=9.5)
        card(s, x, 5.0, w, 1.45, fill=K.PANEL, shadow=False)
        text(s, x + .15, 5.1, w - .3, .25, 'ACTIVITY 5', 9.5, ACC if i < 3 else TEAL, bold=True, autofit=False)
        text(s, x + .15, 5.36, w - .3, 1.0, ex, 11.5, LIGHT, minpt=9)
    tag(D, 'CONCEPT', 'Redrawn from trainer-supplied screenshot (v3.5); ' + HK)

def decomposition(D):
    s = D.new(dark=True); header(D, s, 'Topic 3 · Task decomposition', 'Break the Goal Before You Hand It Out',
                                 'Task decomposition turns one big objective into smaller, specialised jobs that a subagent can finish on its own.')
    card(s, .6, 2.85, 3.5, 2.75, fill=ACC_D, shadow=False)
    text(s, .85, 3.05, 3.0, .28, 'ONE LARGE OBJECTIVE', 11, LIGHT, bold=True, autofit=False)
    text(s, .85, 3.4, 3.0, 2.0, 'A 20-second launch video for the Gula Melaka Cold Brew', 21, WHITE, bold=True, font=SERIF, minpt=14, lh=1.12)
    subs = [('search', 'Gather the launch facts', 'video-researcher'), ('pen', 'Write a 4-scene script', 'scriptwriter'),
            ('megaphone', 'Record the narration', 'voice-over'), ('film', 'Render the 1080×1920 MP4', 'video-producer'),
            ('eyecheck', 'Check it against the QA list', 'video-reviewer')]
    for i, (ic, t, who) in enumerate(subs):
        y = 2.4 + i * .8
        line(s, 4.1, 4.22, 4.85, y + .31, SOFT, 1.5)
        card(s, 4.85, y, 4.3, .62, fill=K.PANEL, shadow=False)
        icon(s, ic, 4.95, y + .08, .46, fill=TEAL)
        text(s, 5.52, y + .06, 2.2, .5, t, 12.5, WHITE, bold=True, anchor='m', minpt=9.5)
        text(s, 7.7, y + .06, 1.4, .5, who, 10.5, ACC, anchor='m', minpt=8)
    text(s, 9.6, 2.42, 3.1, .28, 'A GOOD SPLIT', 11, ACC, bold=True, autofit=False)
    for i, (t, d) in enumerate([('Understands the request', 'Restate the goal before splitting it.'),
                                ('Finds logical subtasks', 'Pieces that can be done separately.'),
                                ('Assigns each an owner', 'One specialist per piece.'),
                                ('Fits back together', 'Each output feeds the next step.')]):
        y = 2.8 + i * .9
        text(s, 9.6, y, .5, .4, f'{i + 1}', 20, ACC, bold=True, font=SERIF, autofit=False)
        text(s, 10.05, y, 2.65, .34, t, 13, WHITE, bold=True, font=SERIF, minpt=10)
        text(s, 10.05, y + .34, 2.65, .5, d, 11, SOFT, minpt=9)
    tag(D, 'CONCEPT', 'Redrawn from trainer-supplied screenshot (v3.5)')

def framework(D):
    s = D.new(dark=True); header(D, s, 'Topic 3 · Task decomposition', 'Five Questions Before Any Agent Starts Work',
                                 'The 5-step decomposition framework. Ask the questions in order: your answers become the prompt that builds the team.')
    from pptx.enum.shapes import MSO_SHAPE as _S
    steps = [(ACC, 'user', 'Understand the primary objective', 'What is the user actually asking for?', 'A 20-second cold brew launch video you can post on 1 Dec.'),
             (TEAL, 'search', 'Identify independent subtasks', 'Which parts can be separated cleanly?', 'Research, script, voice-over, render and review.'),
             (PLUM, 'users', 'Assign specialisation', 'Which subagent suits each task best?', 'Researcher, scriptwriter, voice-over, producer, reviewer.'),
             (OLIVE, 'file', 'Define expected outputs', 'What should each subagent return?', 'research.md · script.md · narration · MP4 · review notes.'),
             (ACC_D, 'clipboard', 'Plan aggregation', 'How will the coordinator combine the results?', 'Mia follows the board; your two HUMAN cards close each stage.')]
    w = 2.36
    for i, (col, ic, t, q, kk) in enumerate(steps):
        x = .6 + i * (w + .07)
        box(s, x, 2.4, w, .95, fill=col, shape=_S.CHEVRON)
        text(s, x + .38, 2.55, .7, .65, f'0{i + 1}', 24, WHITE, bold=True, font=SERIF, anchor='m', autofit=False)
        s.shapes.add_picture(K.icon_png(ic, WHITE), Inches(x + 1.15), Inches(2.62), Inches(.5), Inches(.5))
        text(s, x + .05, 3.5, w - .15, .62, t, 14, WHITE, bold=True, font=SERIF, minpt=11)
        text(s, x + .05, 4.15, w - .15, .62, q, 12, SOFT, italic=True, minpt=9.5)
        card(s, x, 4.95, w - .1, 1.5, fill=K.PANEL, shadow=False)
        text(s, x + .15, 5.05, w - .4, .25, 'KOPI KAKI', 9.5, ACC, bold=True, autofit=False)
        text(s, x + .15, 5.32, w - .4, 1.05, kk, 11.5, LIGHT, minpt=9)
    tag(D, 'CONCEPT', 'Redrawn from trainer-supplied screenshot (v3.5)')

def subagents(D):
    s = D.new(dark=True); header(D, s, 'Topic 3 · Subagents', 'Subagents Are Hired for One Job, Then Released',
                                 'A coordinator spawns a subagent through a task tool. In Hermes that tool is delegate_task.')
    nodes = [('robot', 'Coordinator', 'Mia writes the goal and every fact the helper needs.'),
             ('tool', 'Task tool · delegate_task', 'Starts the helper with that goal and context: up to 10 at once by default.'),
             ('user', 'Subagent', 'Starts with a blank conversation and Mia’s tools; returns only a summary.')]
    for i, (ic, t, d) in enumerate(nodes):
        cx = 2.4 + i * 4.27
        icon(s, ic, cx - .6, 2.35, 1.2, fill=ACC if i != 1 else TEAL)
        if i < 2: line(s, cx + .8, 2.95, cx + 4.27 - .8, 2.95, LIGHT, 2.5, arrow=True)
        text(s, cx - 1.9, 3.65, 3.8, .38, t, 16, WHITE, bold=True, font=SERIF, align='c', autofit=False)
        text(s, cx - 1.8, 4.05, 3.6, .62, d, 12, SOFT, align='c', minpt=9.5)
    text(s, .6, 4.72, 12.1, .32, 'Subagents are spawned on demand. Your named bots are hired once.', 13, LIGHT, italic=True, align='c', autofit=False)
    for i, (col, lab, t, d) in enumerate([(TEAL, 'SUBAGENT', 'Spawned for one job', 'Blank memory · Mia’s tools · reports a summary · gone when the job ends.'),
                                          (ACC, 'NAMED BOT', 'Rafi, Lin, the video team', 'Hired once · keeps its role, memory and chat · you @mention it again tomorrow.')]):
        x = .6 + i * 6.15
        card(s, x, 5.2, 5.95, 1.25, fill=K.PANEL, shadow=False); box(s, x, 5.2, .09, 1.25, fill=col)
        text(s, x + .3, 5.3, 1.6, .28, lab, 10.5, col, bold=True, autofit=False)
        text(s, x + 1.9, 5.27, 3.9, .34, t, 14, WHITE, bold=True, font=SERIF, minpt=11)
        text(s, x + .3, 5.7, 5.5, .65, d, 12, SOFT, minpt=9.5)
    tag(D, 'ARCH', HD + ' ; redrawn from trainer-supplied screenshot (v3.5)')
