"""Visual Day 1: history -> collaboration -> always-on -> Big Five -> real setup.
All diagrams are editable teaching models, not simulated execution screenshots.
"""
from ccdeck import *
from slides_2026 import trace,proof,tag,case_pair
from labs_v23 import NEWS,IOS,ANDROID
ARTICLE='https://buldrr.com/claude-duet-vs-muse-vs-grokbot-vs-dots'
METAHELP='https://www.meta.com/help/superintelligence/1687253048996149/'
DUET='https://claude.com/resources/articles/claude-duet'
GROK='https://docs.x.ai/grok-bot/computer-and-apps'
DOTS='https://learn.chatgpt.com/docs/dots/getting-started'
MANUS='https://manus.im/blog/introducing-manus-2-0'
CUE='https://cue.im/'
HERMES='https://hermes-agent.nousresearch.com/docs/user-guide/desktop'
def source(D,url,label='Official source · checked 29 Oct 2026'):
 s=D.prs.slides[-1];t=text(s,.7,6.72,11.9,.24,label,10,TEAL);t.text_frame.paragraphs[0].runs[0].hyperlink.address=url

def links(D,items,y=6.82,step=3.05):
 s=D.prs.slides[-1]
 for k,(label,url) in enumerate(items):
  t=text(s,.7+k*step,y,step-.1,.24,label,10,TEAL);t.text_frame.paragraphs[0].runs[0].hyperlink.address=url

import os
T1A=os.path.join(os.path.dirname(os.path.abspath(__file__)),'..','..','..','..','courseware','assets','topic1')
def t1a(f):return os.path.normpath(os.path.join(T1A,f))
DOTS_VIDEO='https://www.youtube.com/shorts/gdHwOn2kvT8'
MUSE_VIDEO='https://www.youtube.com/shorts/zuAAjV5dNMA'
GROK_VIDEO='https://www.youtube.com/shorts/Gm_dlNAVOGI'
CUE_VIDEO='https://www.youtube.com/shorts/ZW7gCKTLtpo'
DUET_VIDEO='https://www.youtube.com/shorts/7G5hOD6oF4M'

def _narrow_header(s,w=10.55):
 """Keep kicker/title/subtitle clear of a header-right image."""
 for sh in s.shapes:
  if sh.has_text_frame and sh.top<Inches(2.2) and sh.width>Inches(11): sh.width=Inches(w)

def online_video(s,url,poster,x,y,w,h,name='Online video'):
 """Embed a YouTube video as a PowerPoint online video (Insert > Online Video): it plays inside the slide when clicked.
 The poster frame is what PDF exports and offline viewers show."""
 from pptx.oxml.shapes.picture import CT_Picture
 from pptx.oxml.ns import qn
 from pptx.opc.constants import RELATIONSHIP_TYPE as RT
 vid=url.rstrip('/').split('/')[-1].split('?')[0];embed=f'https://www.youtube.com/embed/{vid}?feature=oembed'
 part=s.part
 v_rId=part.relate_to(embed,RT.VIDEO,is_external=True)
 m_rId=part.relate_to(embed,'http://schemas.microsoft.com/office/2007/relationships/media',is_external=True)
 _,p_rId=part.get_or_add_image_part(poster)
 pic=CT_Picture.new_video_pic(s.shapes._next_shape_id,name,v_rId,m_rId,p_rId,Inches(x),Inches(y),Inches(w),Inches(h))
 for el in pic.iter('{http://schemas.microsoft.com/office/powerpoint/2010/main}media'):
  el.attrib.pop(qn('r:embed'));el.set(qn('r:link'),m_rId)
 s.shapes._spTree.append(pic);s.shapes._add_video_timing(pic)
 return pic

def brief(D,name,sub_,feats,steps,tipt,refs,logo=None,step=3.05,video=None):
 """One briefing slide per always-on agent: what it can do | how to set it up | embedded setup video; app logo top right."""
 s=D.new();header(D,s,'Topic 1 · '+name,name+': What It Does and How to Set It Up',sub_)
 if logo:_narrow_header(s);s.shapes.add_picture(t1a(logo),Inches(11.42),Inches(.5),Inches(1.28),Inches(1.28))
 top,bot=2.3,6.32
 if video:   # three columns: features 4.55 | setup 4.95 | 9:16 player 2.2
  lw,rx,rw=4.55,5.35,4.95;vw=2.2;vh=min(bot-top,vw*16/9);vx=12.7-vw
  card(s,vx-.06,top-.06,vw+.12,vh+.12,fill=INK)
  online_video(s,video[0],t1a(video[1]),vx,top,vw,vh,name=name+' video')
 else:lw,rx,rw=5.6,6.4,6.3
 card(s,.6,top,lw,bot-top);text(s,.85,top+.14,lw-.5,.36,'What it can do',15,ACC,bold=True,font=SERIF,autofit=False)
 rh=(bot-top-.62)/len(feats)
 for i,(ic,t,d) in enumerate(feats):
  y=top+.6+i*rh;icon(s,ic,.85,y+.03,.42)
  text(s,1.42,y-.03,lw-.97,.34,t,13.5,INK,bold=True,font=SERIF,minpt=10.5)
  text(s,1.42,y+.29,lw-.97,rh-.3,d,11 if not video else 10.5,MUTED,minpt=9)
 card(s,rx,top,rw,bot-top,fill=CHAR);text(s,rx+.25,top+.14,rw-.5,.36,'How to set it up',15,ACC,bold=True,font=SERIF,autofit=False)
 rh=(bot-top-.62)/len(steps)
 for i,(t,d) in enumerate(steps):
  y=top+.6+i*rh;circle(s,rx+.25,y+.02,.4,ACC)
  text(s,rx+.25,y+.02,.4,.4,str(i+1),13,WHITE,bold=True,font=SERIF,align='c',anchor='m',autofit=False)
  text(s,rx+.8,y-.03,rw-1.0,.32,t,13.5,WHITE,bold=True,font=SERIF,minpt=10.5)
  text(s,rx+.8,y+.28,rw-1.0,rh-.3,d,11.5 if not video else 10.5,SOFT,minpt=9)
 tip(D,s,tipt,y=6.38);tag(D,'CASE',' ; '.join(u for _,u in refs));links(D,refs,step=step)

def pattern(D,n,title,sub_,rows_,img,shot=None,shot_cap=None,tipt=None):
 """One collaboration pattern: explanation rows (left), the pattern illustration and an optional real screenshot (right)."""
 s=D.new();header(D,s,'Topic 1 · Five collaboration patterns',f'Pattern {n} · {title}',sub_)
 top,bot=2.3,6.3
 lw=4.95 if shot else 12.1-(bot-top)-.25
 card(s,.6,top,lw,bot-top);rh=(bot-top-.2)/len(rows_)
 for i,(ic,t,d) in enumerate(rows_):
  y=top+.14+i*rh;icon(s,ic,.82,y+.04,.4)
  text(s,1.36,y-.02,lw-.95,.32,t,13,ACC if t.startswith('Real') else INK,bold=True,font=SERIF,minpt=10.5,autofit=False)
  text(s,1.36,y+.28,lw-.95,rh-.32,d,11,MUTED,minpt=9)
 side=bot-top
 picture(s,t1a(img),12.7-side,top,side,side,frame=True)
 if shot:
  x0=.6+lw+.22;bw=12.7-side-.22-x0
  text(s,x0,top-.02,bw,.26,shot_cap or 'Real example',10,TEAL,bold=True,autofit=False,align='c')
  picture(s,t1a(shot),x0,top+.32,frame=True,maxw=bw,maxh=side-.36)
 if tipt:tip(D,s,tipt,y=6.42)
 return s

GENIC='https://www.genicsolutions.com/ai-agents-for-business-use-cases-and-implementation/'
AGENTSEA='https://www.straitstimes.com/tech/new-platform-to-create-ai-agents-available-to-all-public-healthcare-professionals'
def agentsea(D):
 """Singapore news: Synapxe AgentSea (Straits Times, 24 Aug 2026), straight after the AI history."""
 s=D.new();header(D,s,'Topic 1 · Singapore in the news','Singapore Healthcare: 12,000+ AI Agents in Three Months',
  'AgentSea, Synapxe’s agentic AI platform built with AWS, lets public healthcare staff build their own agents in plain language')
 for j,(big,lab,de) in enumerate([('12,000+','AI agents created','By healthcare professionals, from the end-May 2026 launch to 24 Aug 2026.'),
   ('80,000+','healthcare professionals','Everyone in Singapore’s public health sector can create agents, with little or no coding.'),
   ('300+','agents shared','Published for colleagues to reuse and adapt for their own work.')]):
  x=.6+j*4.1;card(s,x,2.3,3.9,1.95,fill=CHAR if j==0 else WHITE)
  text(s,x+.25,2.36,3.4,.85,big,40,ACC,bold=True,font=SERIF,autofit=False)
  text(s,x+.25,3.18,3.4,.32,lab,14,WHITE if j==0 else INK,bold=True,font=SERIF,minpt=11,autofit=False)
  text(s,x+.25,3.5,3.45,.7,de,11,SOFT if j==0 else MUTED,minpt=9)
 for j,(ic,ti,de) in enumerate([('user','Cardiologist, National Heart Centre','An agent summarises patient histories before clinic, halving 40+ minutes of daily pre-clerking. He still verifies and makes the clinical assessment.'),
   ('calendar','KK Women’s and Children’s Hospital','An agent generates daily staff rosters, cutting manual data entry and transcription errors.'),
   ('shield','Guardrails and governance','Rejects requests for card or ID numbers; MOH’s updated AI in Healthcare Guidelines keep professional judgment with people.')]):
  x=.6+j*4.1;card(s,x,4.45,3.9,1.8);icon(s,ic,x+.22,4.6,.42)
  text(s,x+.75,4.6,3.0,.42,ti,12.5,INK,bold=True,font=SERIF,minpt=10,anchor='m')
  text(s,x+.22,5.12,3.5,1.08,de,11,MUTED,minpt=9)
 tip(D,s,'The course pattern at national scale: agents prepare the work; clinicians verify and decide.',y=6.38)
 tag(D,'CASE',AGENTSEA);links(D,[('The Straits Times, 24 Aug 2026: New platform to create AI agents available to all public healthcare professionals · click',AGENTSEA)],step=12)
 return s

SCI_NEWS=[('Quanta Magazine · 8 Sep 2026','https://www.quantamagazine.org/ai-has-solved-one-of-maths-1-million-millennium-prize-problems-20260908/'),
 ('Stanford Medicine · 17 Sep 2026','https://med.stanford.edu/news/all-news/2026/09/virtual-biotech-company.html'),
 ('Nature · 7 Oct 2026','https://www.nature.com/articles/d41586-026-03196-8'),
 ('Firstpost · 27 Oct 2026','https://www.firstpost.com/tech/openai-claims-three-new-ai-breakthroughs-in-math-physics-and-computer-science-14055633.html')]
def science_news(D):
 """2026 AI-for-science news: four results, each with its verification status."""
 s=D.new();header(D,s,'Topic 1 · AI for Science in the news','2026 Breakthroughs: Agents at Scale, Humans Verify',
  'Four reported results. Each shows what the agents did, and what still has to be checked before it counts as a discovery.')
 items=[('flask','Maths · a Millennium Prize problem','OpenAI ran about 10,000 autonomous agents to show the 3D Navier–Stokes fluid equations can “blow up”.','Checked in the Lean proof assistant; humans confirmed the formal statement. Priority dispute with two mathematicians.'),
  ('users','Drug discovery · a virtual biotech','Stanford’s 37,000 AI agents catalogued about 50,000 clinical trials in under a week and designed a lung-cancer antibody–drug conjugate.','A pharma company independently reached the same design; real-lab testing of other findings is next.'),
  ('file','Maths at volume · 700+ preprints','OpenAI posted more than 700 AI-generated maths preprints from an unreleased model.','Mathematicians call them a “messy first draft” and raise credit and pre-emption concerns.'),
  ('target','Maths, physics and computing claims','OpenAI claims proofs on Kakeya conjectures and 4D quantum field theory, and an algorithm that would disprove the Exponential Time Hypothesis.','Claims awaiting verification; expert panels are reviewing them.')]
 for j,(ic,ti,de,chk) in enumerate(items):
  x=.6+(j%2)*6.15;y=2.25+(j//2)*2.0;card(s,x,y,5.95,1.88)
  icon(s,ic,x+.2,y+.16,.42);text(s,x+.75,y+.12,4.95,.32,ti,13.5,INK,bold=True,font=SERIF,minpt=11,autofit=False)
  text(s,x+.75,y+.42,5.0,.36,SCI_NEWS[j][0],10,TEAL,bold=True,autofit=False)
  text(s,x+.2,y+.74,5.55,.56,de,11,MUTED,minpt=9)
  rich(s,x+.2,y+1.3,5.55,.55,[[('CHECK  ',dict(color=ACC,bold=True,size=10.5)),(chk,dict(color=INK,size=10.5))]])
 tip(D,s,'Agents now propose at massive scale. Proof checkers, laboratories and experts still decide what counts as a discovery.',y=6.3)
 tag(D,'CASE',' ; '.join(u for _,u in SCI_NEWS));links(D,[(t+' · click',u) for t,u in SCI_NEWS],step=3.05)
 return s

def analogy(D):
 """Visual analogy: Human = Brain + Body --food--> Work; AI agent = Model + Harness --tokens--> Work."""
 s=D.new();header(D,s,'Topic 1 · Under the hood','An AI Agent Works Like a Person','Brain + body need food to do work. Model + harness need tokens to do work.')
 rows_=[('HUMAN','user',PLUM,[('brain','Brain','Thinks, plans, decides'),('hand','Body','Hands, eyes and tools act')],('coffee','Food','Energy'),('briefcase','Work','The job gets done')),
        ('AI AGENT','robot',TEAL,[('sparkles','Model','Reasons, proposes next step'),('gear','Harness','Tools, memory, skills, controls')],('coins','Tokens','Fuel and cost'),('check','Work','Verified output'))]
 def node(x,ic,t,d,fill):
  circle(s,x-.55,cy-.55,1.1,fill);icon(s,ic,x-.42,cy-.42,.84,fill=fill)
  text(s,x-1.15,cy+.62,2.3,.4,t,19,INK,bold=True,font=SERIF,align='c',autofit=False)
  text(s,x-1.15,cy+.98,2.3,.36,d,11,MUTED,align='c',minpt=9)
 for r,(lab,lic,col,parts,fuel,work) in enumerate(rows_):
  y=2.2+r*2.08;cy=y+.72
  card(s,.6,y,12.1,1.95,fill=WHITE);box(s,.6,y,.12,1.95,fill=col)
  circle(s,1.0,cy-.42,.84,col);icon(s,lic,1.1,cy-.32,.64,fill=col)
  text(s,.75,cy+.5,1.35,.3,lab,12,col,bold=True,align='c',autofit=False)
  node(3.25,*parts[0],ACC);text(s,4.45,cy-.42,.6,.84,'+',40,ACC,bold=True,font=SERIF,align='c',anchor='m',autofit=False);node(6.25,*parts[1],ACC)
  box(s,7.4,cy-.08,2.15,.16,fill=col);box(s,9.45,cy-.3,.45,.6,fill=col,shape=MSO_SHAPE.ISOSCELES_TRIANGLE).rotation=90
  box(s,7.6,cy+.22,1.75,.62,fill=col,shape=MSO_SHAPE.ROUNDED_RECTANGLE)
  icon(s,fuel[0],7.66,cy+.28,.5,fill=col);text(s,8.18,cy+.22,1.15,.36,fuel[1],15,WHITE,bold=True,font=SERIF,autofit=False)
  text(s,8.18,cy+.52,1.15,.28,fuel[2],9.5,WHITE,autofit=False)
  node(10.85,*work,CHAR)
 tip(D,s,'No food, no work; no tokens, no work. Every read, reasoning step and reply spends tokens, so a token budget is the agent’s running cost.',y=6.4)
 tag(D,'MODEL','Trainer analogy: human brain/body/food vs model/harness/tokens')
 return s

def illustrated(D,k,title,steps,img,tipt,source):
 """Four numbered steps (left) beside a wide illustration (right)."""
 s=D.new();header(D,s,k,title)
 top,bot=1.95,6.15;card(s,.6,top,5.25,bot-top);rh=(bot-top-.2)/len(steps)
 for i,(ic,t,d) in enumerate(steps):
  y=top+.14+i*rh;circle(s,.82,y+.05,.42,ACC)
  text(s,.82,y+.05,.42,.42,str(i+1),13,WHITE,bold=True,font=SERIF,align='c',anchor='m',autofit=False)
  if i<len(steps)-1:line(s,1.03,y+.5,1.03,y+rh+.02,color=SOFT,w=1.25)
  text(s,1.42,y-.02,4.25,.34,t,14,INK,bold=True,font=SERIF,minpt=11,autofit=False)
  text(s,1.42,y+.31,4.25,rh-.34,d,11.5,MUTED,minpt=9)
 picture(s,t1a(img),6.1,top,frame=True,maxw=6.6,maxh=bot-top)
 tip(D,s,tipt,y=6.32);tag(D,'ARCH',source);return s

def topic1(D,C,mark):
 """Topic 1 narration: what agents are -> how they work -> where people decide -> where they are used
 -> the always-on agents (one briefing slide each) -> Activity 1 checks -> Lab 1 (the topic ends with its lab)."""
 mark('t1a')
 section(D,'Topic 1',C.TOPICS[0]['title'],['Brief history and what an agent is','Under the hood: how agents work','AI does the work; humans decide','Five ways people and agents work together','Use cases and the Big Five always-on agents','Activity 1 · Set up one assistant'])
 timeline(D,'Topic 1 · Brief history','From Rules to Learning to Action',[
 ('Rules','1950s–80s','Explicit rules and expert systems.'),('Learning','1990s–2010s','Patterns learned from examples.'),('Language','2017–2023','Transformers and chat interfaces.'),('Action','2024 onward','Tool use, feedback and goal-directed work.')],foot='These approaches coexist. A modern agent still uses ordinary software, rules and databases.');tag(D,'MODEL','Attention Is All You Need 2017; ReAct 2022')
 timeline(D,'Topic 1 · Brief history','2023 → 2026: Four Engineering Lenses',[
 ('Prompt Engineering','2023','Specify goal, audience and output.'),('Context Engineering','2024','Provide relevant facts and instructions.'),('Harness Engineering','2025','Connect tools, state and checks.'),('AI Agents','2026','Delegate outcomes and inspect results.')],foot='Course framing: overlapping priorities, not invention dates or exclusive eras.');tag(D,'MODEL','Trainer-requested teaching sequence; agent and context engineering literature')
 agentsea(D)
 compare(D,'Topic 1 · Agents and agentic AI','Agentic AI Is the System. An Agent Does the Work.',
 ('Agentic AI','network',['Goal-directed planning, action and feedback.','May contain one agent or a coordinated team.','The harness controls access, state and execution.']),
 ('Always-on agent · course focus','clock',['A persistent, 24/7-capable agentic system.','Can act between messages while runtime is available.','Examples: OpenClaw, Hermes, Muse and Dots.']),tipt='Always-on describes the deployment we study; it is not a universal definition of every AI agent or an uptime promise.');tag(D,'ARCH','Anthropic building effective agents; vendor agent documentation')
 # Under the hood (before Lab 1): analogy -> the model inside its harness -> tools and skills -> the loop.
 analogy(D)
 hub(D,'Topic 1 · Under the hood','What Makes an AI Agent: a Model Inside a Working System','AI agent',[
 ('brain','Model','Reasons and proposes the next action.'),('tool','Tools','Read, calculate, browse, render.'),('book','Skills','Reusable procedures and checks.'),('memory','Memory','Approved persistent facts.'),('folder','Knowledge base','Curated evidence and versions.'),('file','Instructions','Goal, policy and output contract.'),('shield','Runtime controls','Loop, permissions, limits and logs.')],foot='A more capable model cannot compensate for missing access, poor sources or an unverified action.');tag(D,'ARCH',HERMES)
 compare(D,'Topic 1 · Under the hood','Tools Execute. Skills Explain How.',
 ('Tool','tool',['gog reads a calendar through a real API.','Python calculates a workbook total.','FFmpeg renders an actual MP4.']),
 ('Skill','book',['gog-office sets read limits and source checks.','professional-ppt defines a repeatable slide process.','launch-video defines scenes, render and QA.']),tipt='Installing a skill creates neither the executable nor Google permission. Verify all three separately.');tag(D,'ARCH',HERMES)
 loop(D,'Topic 1 · Under the hood','The Agentic Loop: Observe, Act, Verify',[
 ('target','Goal','Bounded outcome'),('search','Observe','Read permitted evidence'),('brain','Select','Choose next action'),('tool','Act','Execute tool'),('eyecheck','Verify','Inspect real result'),('hand','Decide','Repeat, ask or stop')],foot='Tool errors are observations. Never replace an unavailable result with an invented success.');tag(D,'ARCH','ReAct; Hermes documentation')
 # Where people stay in charge.
 compare(D,'Topic 1 · Human–AI work','AI Agents Do the Work. Humans Decide.',
 ('AI agent','robot',['Reads the permitted source records.','Uses tools to calculate, draft and produce.','Checks the result and reports exceptions.']),
 ('Human','user',['Chooses the goal and permitted access.','Exercises judgment when facts or values conflict.','Approves consequential actions and accepts work.']),tipt='Example: agent proposes a meeting time; you resolve a conflict and approve the message.');tag(D,'MODEL','IBM human-AI collaboration; classroom task contract')
 compare(D,'Topic 1 · Two collaboration modes','Personal Assistant or Digital Worker?',
 ('Personal assistant','calendar',['Organises YOUR day: email, schedule, travel.','Success: a source-checked brief and useful draft.','You decide priorities and approve commitments.']),
 ('Digital worker','briefcase',['Performs a BUSINESS role: analyst, coder, bookkeeper.','Success: verified work products and a review trail.','A named owner accepts the result and exceptions.']),tipt='These are work roles, not two different kinds of language model.');tag(D,'MODEL','IBM; BULDRR job-fit framing; course design')
 # Five collaboration patterns: from one person and one agent to agents working across companies.
 pattern(D,1,'One Human, One AI Agent','A personal assistant: one owner, one agent, one set of permissions',[
  ('user','How it works','The agent watches your mail and calendar, then messages you with reminders, clashes and drafts.'),
  ('phone','Real example','A hosted assistant flags a weekend clash: two booked course days and a new teaching invitation.'),
  ('hand','The human decides','Which commitment wins, what to reply, and every send or acceptance.'),
  ('alert','Watch for','The wrong account or stale facts. Open the linked email before acting on a reminder.')],
  'pattern1.jpg','p1-assistant-chat.png','Real example · names blurred',tipt='Simplest pattern and the one in Lab 1: ownership is clear, but quality depends on one person reviewing.')
 pattern(D,2,'Multiple Humans, One AI Agent','A shared team agent in a group chat serves several people',[
  ('users','How it works','Team members ask one agent in a shared channel; it updates systems and reports back to the group.'),
  ('chat','Real example','In a trainers’ group chat, an OpenClaw agent reassigns a course run in the LMS and reports a blocker.'),
  ('hand','The humans decide','Who may instruct the agent, whose request wins in a conflict, and who confirms the change.'),
  ('alert','Watch for','Conflicting instructions and over-broad access. Log who asked for what.')],
  'pattern2.jpg','p2-group-chat.png','Real example · names blurred',tipt='A good shared agent reports what it verified and what is still unknown, and escalates what it cannot fix.')
 pattern(D,3,'One Human, Multiple AI Agents','One owner directs a team of specialist agents',[
  ('robot','How it works','Each agent is a separate Hermes profile with its own role, model, skills and running gateway.'),
  ('server','Real example','Six profiles form a YouTube line: research → script → video → optimise → publish, plus the default agent.'),
  ('hand','The human decides','Splits the goal, sets hand-offs and approval gates, and accepts the final output.'),
  ('alert','Watch for','Lost hand-offs, duplicated work and cost per agent. Use a Kanban board and an auditor.')],
  'pattern3.jpg','p3-hermes-profiles.png','Real example · Hermes profiles',tipt='This is the Day 2 pattern: your Hermes Kanban video team in Lab 5 and Lab 8.')
 pattern(D,4,'Multiple Humans, Multiple AI Agents','A department: several owners, each with agents, sharing one board',[
  ('users','How it works','People and agents share one task board; every card has a human owner, an agent assignee and a status.'),
  ('briefcase','Example','Marketing, finance and operations staff each run agents that research, calculate and draft one launch plan.'),
  ('hand','The humans decide','Each owner approves their own agents’ work; a lead resolves priorities across teams.'),
  ('alert','Watch for','Unclear accountability, agents overwriting each other’s files, and access creeping across teams.')],
  'pattern4.jpg',tipt='Scale needs governance: named owners, a shared board, permissions per role and an audit trail.')
 pattern(D,5,'Inter-Company Collaboration','Agents from different organisations exchange work across a company boundary',[
  ('send','How it works','Company A’s agent sends a request, quote or data package to Company B’s agent; each side runs its own checks.'),
  ('briefcase','Example','A buyer’s procurement agent asks for quotes; the supplier’s sales agent checks stock and replies with an offer for sign-off.'),
  ('hand','The humans decide','Contracts, prices, what data may be shared, and final acceptance on both sides.'),
  ('alert','Watch for','Who the other agent really is, personal data (PDPA), injected instructions and liability for commitments.')],
  'pattern5.jpg',tipt='Emerging pattern: treat everything from another company’s agent as untrusted data until a human verifies it.')
 # Five business use cases (Genic Solutions, 30 Sep 2026); AI for Science and the humanoid task follow.
 cases=[
 ('Customer service operations','Incoming requests + account history','Classify, look up the account, draft a reply, suggest next step','Complex or sensitive cases; limits on what it may complete','Prepared response + next action','Less time spent searching for information','Digital worker · customer service'),
 ('IT service management','Access, account, device and software requests','Collect details, check known issues, create and route tickets','Technician fixes; approves access changes','Triage-ready ticket with details','Repeat requests routed to the right group','Digital worker · IT support'),
 ('Finance operations','Invoices, payment queries, internal records','Read documents, match records, flag missing details','Approves payments and high-risk actions','Items prepared for approval','Clear approval rules applied every time','Digital worker · finance'),
 ('Sales operations','Customer records + meeting notes','Update records, prepare notes, create follow-up tasks','Salesperson owns the customer relationship','Updated records + follow-up tasks','Less administration for sales staff','Digital worker · sales'),
 ('Procurement and supply chain','Orders, stock, supplier and delivery data','Monitor orders, flag delays, gather facts across systems','Operations manager decides the response','Delay alert with consolidated records','Delays flagged with supporting records','Digital worker · operations')]
 for i in range(0,len(cases),2):
  case_pair(D,'AI Agents for Business: Five Use Cases','handshake',cases[i:i+2],i+1,kick='Topic 1 · Use cases',sub='Business functions · what the agent does, what people decide, and the evidence to check')
  links(D,[('Genic Solutions, 30 Sep 2026: AI Agents for Business, Use Cases and Implementation · click',GENIC)],step=12);tag(D,'CASE',GENIC)
 illustrated(D,'Topic 1 · Scientific collaboration','AI for Science: Proposal → Experiment → Discovery',[
 ('user','Research goal','Scientist sets the question and constraints.'),('users','Hypotheses','An AI co-scientist proposes, critiques and ranks ideas.'),('flask','Experiment','Scientists test the selected hypotheses in the lab.'),('check','Validation','New evidence supports or rejects the claim.')],
 'ai-for-science.jpg','A plausible proposal is not a discovery. Record provenance and experimental results.','https://research.google/blog/accelerating-scientific-breakthroughs-with-an-ai-co-scientist/')
 science_news(D)
 illustrated(D,'Topic 1 · Future collaboration','Humanoid + Human: A Bounded Physical Task',[
 ('eye','Sense','The robot observes the authorised room and the resident it is helping.'),('brain','Plan','It proposes one bounded task, such as handing over water or a walking aid.'),('user','Authorise','A caregiver confirms the task and keeps the stop control.'),('check','Verify','The caregiver checks the hand-off, the resident’s response and recovery behaviour.')],
 'humanoid-elderly-care.jpg','Prospective design for elderly care. Manufacturer demos are not proof of safe deployment in homes or care facilities.','https://deepmind.google/models/gemini-robotics/')
 mark('t1b')
 timeline(D,'Topic 1 · Deployment evolution','OpenClaw → Hermes → Hosted Personal Agents',[
 ('OpenClaw','Configurable harness','Local/server tools and channels.'),('Hermes Agent','Desktop + skills','Local workflow, memory and reusable procedures.'),('Hosted agents','Managed cloud','Dots, Muse, Duet, Grok Bot and Cue.'),('Your choice','Work + control','Access, oversight, cost and data boundaries.')],foot='A teaching progression in deployment choices, not corporate lineage or a claim one product replaced another.');tag(D,'ARCH','Official OpenClaw, Hermes, Meta, Anthropic, xAI and OpenAI docs')
 table(D,'Topic 1 · The Big Five','The Big Five: Same Always-On Idea, Different Jobs',
 ['Product','Useful starting job','Human control to inspect','Access checkpoint'],[
 ('Muse','Schedule, errands and personal email','Connector permissions + sensitive-action approval','Singapore since 27 Oct; free core use'),
 ('Dots','Persistent goals and delegated tasks','Connection permissions, custom rules and review','Check access and preview limits for your account'),
 ('Grok Bot','Named workers sharing ongoing work','Approval settings; shared cloud logins across bots','Check your plan’s access'),
 ('Duet','Work briefs and business follow-ups','Approval before sensitive actions; some actions disallowed','US Pro/Max waitlist preview'),
 ('Cue','Personal errands, bookings and calls','Per-agent permissions, wallet budget, review boundaries','Early access; free with invite code')],colw=[1.4,3.4,4.6,2.7],size=12.5,sub='Five hosted always-on agents; one briefing slide each follows, with its official source.',tipt='No universal winner. Verify your current account, region, connector and action controls before delegating.');tag(D,'MODEL',ARTICLE+' ; '+DUET+' ; '+NEWS+' ; '+GROK+' ; '+DOTS);source(D,ARTICLE,'Comparison reference: BULDRR · click to read')
 brief(D,'Meta Muse','Meta · available in Singapore since 27 Oct 2026 · recommended route for Activity 1',[
  ('check','Does the work','Manages email, coordinates schedules and turns long-term goals into action plans.'),
  ('plug','Connected apps','Gmail, Google Calendar and Instagram; Singapore partners such as Grab, Shopee and Singapore Airlines are rolling out.'),
  ('shield','Asks before acting','Checks with you before sending an email or making a purchase; retrieve-only connectors available.'),
  ('lock','Secure and mostly free','Muse Secure VM keeps passwords and payment methods out of the agent’s view; iOS, Android and muse.ai.')],
  [('Install','Muse from Meta (iOS / Android) or muse.ai.'),
   ('Sign in','Complete onboarding; link your Meta accounts.'),
   ('Connect Gmail','Settings → Connectors → Connect beside Gmail → Continue.'),
   ('Authorise','Finish the Google login in the browser; repeat for Google Calendar.'),
   ('Approvals','Keep retrieve-only where offered; approve every send or invite.')],
  'Shortcut: ask “Connect my Gmail”. Partner features roll out progressively; no bookings or purchases in class.',
  [('Meta Newsroom: Muse in SG · click',NEWS),('Meta Help: Connectors · click',METAHELP),('Muse iOS (SG) · click',IOS),('Muse Android (SG) · click',ANDROID),('YouTube: 10 ways to use Muse',MUSE_VIDEO)],logo='logo-muse.png',video=(MUSE_VIDEO,'video-muse.jpg'),step=2.44)
 brief(D,'OpenAI Dots','OpenAI · ChatGPT Learn documentation, checked 29 Oct 2026 · check your account’s access',[
  ('target','Ongoing responsibilities','Give one dot several long-running jobs instead of one-off chats.'),
  ('server','Its own cloud computer','Works in its own browser; optionally on your computer while the ChatGPT app is open.'),
  ('puzzle','Plugins and channels','Email, calendar and files through plugins; Slack and other contact methods for messages.'),
  ('clock','Delegates and schedules','Hands parts to ChatGPT Work or Codex; fixed schedules, event triggers and an Activity view.')],
  [('Open dots','ChatGPT mobile app, desktop app or desktop browser (not mobile web).'),
   ('Create a dot','Follow the intro; name it and pick its look (changeable later).'),
   ('Plugins','Connect email, calendar and files plugins; authorise each one.'),
   ('Channels','Optional: profile → Add a contact method such as Slack.'),
   ('Assign work','Describe an ongoing responsibility, its sources and your decisions.')],
  'A messaging channel does not grant inbox access. Ending a chat does not stop assigned work: check Activity.',
  [('ChatGPT Learn: Get started with your dot · click',DOTS),('YouTube: How to set up your first dot (Short) · click',DOTS_VIDEO)],logo='logo-dots.png',video=(DOTS_VIDEO,'video-dots.jpg'),step=4.2)
 brief(D,'Grok Bot','xAI · Grok Bot documentation, updated 28 Oct 2026 · check your plan’s access',[
  ('server','Persistent cloud computer','Browser, command line and files; keeps working while your laptop is closed.'),
  ('users','Several bots, one computer','Each bot has its own screen; sessions, files and logins are shared across your bots.'),
  ('mail','Apps, MCP and its own email','Gmail, Calendar, Sheets, Slack and Notion; mail to name@mail.grokbot.com can start routines.'),
  ('eye','Watch, take over, review','Live computer view; take control for a sign-in or CAPTCHA; Auto Review gates sends.')],
  [('Install','Start at x.ai/bot; get the desktop or mobile app and sign in.'),
   ('Create a bot','Sidebar + → Create a Bot → name it → Create.'),
   ('Connect apps','Open Connect apps (mobile: avatar → Connect apps); add Gmail.'),
   ('Authorise','Finish the Google login in the browser; repeat for Google Calendar.'),
   ('Verify','Both show under Installed; “Added” is not signed in, so fix “Needs auth”.')],
  'All bots on one account share a cloud computer and logins. The bot’s @mail.grokbot.com inbox is not your Gmail.',
  [('xAI docs: Grok Bot computer and apps · click',GROK),('YouTube: 3 steps to set up Grok Bot (Short) · click',GROK_VIDEO)],logo='logo-grok.png',video=(GROK_VIDEO,'video-grok.jpg'),step=4.2)
 brief(D,'Claude Duet','Anthropic research preview · 21 Oct 2026 · US Pro/Max via waitlist · reference route',[
  ('calendar','Morning brief, then watching','A daily brief, then checks every few hours; messages you only when something needs you.'),
  ('server','Works while you are away','Its own virtual machine and browser: meetings, dashboards, forecasts, prototypes and web apps.'),
  ('puzzle','6,000+ connectors','Gmail, Slack, QuickBooks, Stripe, HubSpot, Docusign, Canva, Shopify and more.'),
  ('shield','You approve actions','Shows exactly what it will send, pay, sign or delete and waits; some actions stay off-limits.')],
  [('Get access','Pro or Max plan; join the Claude Duet waitlist on claude.com.'),
   ('Install','Open the Claude Duet app (iOS / Android) or claude.ai once invited.'),
   ('Carry context','Memory, skills and apps carry over from your Claude account.'),
   ('Connect apps','Add Gmail, Calendar, Slack or others from the Claude Marketplace.'),
   ('Set the brief','Tell it what to watch and how often; approve every send or payment.')],
  'Reference route, not assessed. Background checks use your plan limits; some actions stay off-limits even with approval.',
  [('Anthropic: Introducing Claude Duet · click',DUET),('YouTube: Claude Duet is here (Short) · click',DUET_VIDEO)],logo='logo-duet.png',video=(DUET_VIDEO,'video-duet.jpg'),step=4.2)
 brief(D,'Manus Cue','Manus · early access since 28 Sep 2026 · free with an invite code · reference route',[
  ('user','Agent identity','Each agent has its own email address, phone number, wallet and computer.'),
  ('phone','Calls, messages, payments','Takes forwarded calls, emails from its own address and pays within the budget you set.'),
  ('users','Agent teams','Agents share a group chat and a goal; you make the final call.'),
  ('puzzle','Connectors and routines','Gmail, Trip.com, Oura and more; routines repeat your steps; scan a QR code to create an agent.')],
  [('Get Cue','cue.im on the web, desktop or Android; iOS pending.'),
   ('Invite code','Sign up with an invite code from your trainer or a user.'),
   ('Create an agent','Give it one responsibility and the context it needs.'),
   ('Identity + limits','Choose its email, phone, wallet; set budgets and approvals.'),
   ('Connect + routines','Connect Gmail or other services, then define routines for repeat steps.')],
  'Reference route, not assessed. Keep wallet budgets small; require approval for payments, calls and messages.',
  [('Cue by Manus: cue.im · click',CUE),('YouTube: Meet Manus Cue (Short) · click',CUE_VIDEO)],logo='logo-cue.png',video=(CUE_VIDEO,'video-cue.jpg'),step=4.2)

def pp1_evidence(D,kick):
 """PP Task 1 evidence, shown inside Lab 1."""
 proof(D,kick,'PP Task 1 Has Two Images and Written Analysis',[
 ('camera','Image A · morning','Own chosen assistant request/result: real schedule check or email draft.'),('file','Image B · afternoon','Own Hermes-generated artifact: PPT, dashboard, analysis or video.'),('eyecheck','Keep working evidence','Source references, actual artifact and tool/run trace.'),('pen','Written analysis','Explain underlying LLM/agent loop, strengths and limitations for the paper’s client.')],'Current assessment v7; A1 and A3 unchanged')

def office_extra(D):
 units=[
 ('Install a Skill Only After Reviewing Its Source',[('search','Origin','Official or supplied reviewed source; record version.'),('book','Inspect','Read SKILL.md and any scripts it can run.'),('lock','Destination','Selected Desktop profile’s actual Skills folder.'),('check','Probe','Reload, list skill and run a harmless tool test.')]),
 ('Five Skills, One Hermes Agent',[('mail','gog-office','Read sources with bounds and provenance.'),('chart','office-analytics','Schema, formulas and reconciliation.'),('file','professional-ppt','Editable decision story and rendered check.'),('code','interactive-dashboard','Filters, sorting and CSV export, tested offline.'),('film','launch-video','Approved scenes, render and video QA.')]),
 ('Native Desktop and WSL Are Separate Runtimes',[('server','Native app','Uses native paths, PATH and credential store.'),('code','WSL','Linux paths and executables are separate.'),('alert','Symptom','Terminal works, Desktop cannot find gog.'),('check','Fix','Probe inside the selected Desktop runtime; restart only that gateway if needed.')]),
 ('Check Skill, Tool and Permission Separately',[('book','Skill','Does Hermes load the intended SOP?'),('tool','Tool','Does gog/FFmpeg execute in that runtime?'),('lock','Permission','Does the assigned account have the needed service grant?'),('file','Output','Can the real source/artifact be opened independently?')]),
 ('The Office Pack Has a Stable Data Contract',[('file','Brief','outputs/lab01/research-brief.md'),('chart','Analysis','outputs/lab02/analysis.json and analysis.xlsx'),('file','Artifacts','outputs/lab02/dashboard.html and management-brief.pptx'),('film','Videos','Day 1: outputs/single-video; Day 2: outputs/team/video')]),
 ('First Authorise Read-Only Gmail',[('login','Client','Trainer-supplied Desktop OAuth client.'),('user','Account','Assigned Google identity; browser consent.'),('lock','Scope','--services gmail,calendar,drive --gmail-scope readonly'),('shield','Runtime','--readonly and Gmail --gmail-no-send on reads.')]),
 ('One Test Send Is a Separate Procedure',[('pen','Prepare','Exact recipient, subject, body and attachments.'),('user','Approve','Trainer-approved training recipient; explicit SEND decision.'),('lock','Grant','Inspect installed help; preserve existing services when reauthorising write.'),('check','Read back','Send once; verify returned message ID and actual Sent record.')]),
 ('Read-Only Defaults Resume After the Test',[('shield','Default','gog-office remains the read-only SOP.'),('mail','Exception','Only the approved test-send procedure may send.'),('eyecheck','Inspect','One send result; no hidden follow-up automation.'),('check','Restore','Next reads again use --readonly and --gmail-no-send.')]),
 ('An Email Draft Is Not a Sent Message',[('pen','Draft text','Local text has no mailbox delivery state.'),('file','Server draft','Saved draft has an ID but has not been sent.'),('mail','Send result','Real send returns a message ID.'),('check','Verification','Inspect actual Sent record for To/Subject/body.')]),
 ('Make a Single-Agent Video Before a Team Video',[('book','Skill','launch-video supplies the procedure.'),('tool','Renderer','One Hermes agent invokes local FFmpeg tool.'),('film','Output','Real 15s 1080x1920 H.264 MP4.'),('user','Review','Watch full video; approve exact hash.')]),
 ('A Storyboard Is an Input, Not a Completed Video',[('pen','Script','Three approved five-second scenes.'),('code','Render','Run the actual reviewed local tool.'),('eye','Playback','Inspect all scenes and ffprobe metadata.'),('check','Accept','Store MP4, QA evidence and human decision.')]),
 ('Video Rights and Facts Need Checks',[('file','Facts','Approved price, launch date and offer terms.'),('image','Assets','Use supplied or appropriately licensed assets.'),('shield','Claims','No unsupported health/certification promises.'),('user','Decision','Human checks every scene before accepting.')]),
 ('A Small Video Change Invalidates Old Acceptance',[('file','Version 1','Watch, audit and hash actual MP4.'),('user','Revision','Human requests one concrete scene correction.'),('tool','Version 2','Re-render and repeat scene/metadata checks.'),('check','Accept','Log new hash; previous approval does not carry over.')]),
 ('Total Time Includes Setup and Human Review',[('clock','Setup','Installation, consent and source preparation.'),('robot','Run','Agent work and deterministic tool execution.'),('eyecheck','Review','Source checks, render inspection and corrections.'),('chart','Compare','Use total time per accepted artifact, not model speed alone.')]),
 ('Office Failures Should Stay Visible',[('alert','Missing service','Mark Calendar unavailable rather than empty.'),('chart','Mismatch','Stop when aggregates disagree with source rows.'),('file','Bad output','Open PPT/MP4; completion text is insufficient.'),('hand','Human decision','Repair or record incomplete; never manufacture evidence.')]),
 ('Create a Reusable Skill from a Real Correction',[('eyecheck','Observe','Find an actual calculation or layout defect.'),('book','Improve','Add a specific check to the reviewed skill.'),('loop','Repeat','Run a new bounded task from source.'),('check','Evaluate','Compare defect rate, review time and accepted output.')])]
 return {t:(lambda k,t=t,items=items:proof(D,k,t,items,HERMES+' ; https://gogcli.sh/gmail-workflows.html')) for t,items in units}
