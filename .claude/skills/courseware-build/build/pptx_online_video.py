"""Embed YouTube videos in a python-pptx deck so they PLAY INSIDE PowerPoint when clicked.

Reproduces PowerPoint's own Insert > Video > Online Film (Mac) / Online Video (Windows) markup, copied from a
sample the trainer inserted in PowerPoint for Mac and confirmed playing (30 Oct 2026):

  slide rels   <Relationship Type=".../officeDocument/2006/relationships/video"
                 Target="https://www.youtube.com/embed/<id>?feature=oembed" TargetMode="External"/>
  <p:pic>      cNvPr carries <a:hlinkClick r:id="" action="ppaction://media"/>;
               nvPr holds ONLY <a:videoFile r:link="rIdN"/> (no p14:media, no p15:webVideoPr, no media rel);
               blipFill = the poster image.
  <p:timing>   THIS is what makes a click play the video. For each video:
               - mainSeq clickEffect  cmd="playFrom(0.0)"  (presetClass="mediacall", presetID 1)
               - <p:video><p:cMediaNode vol="80000"> node
               - interactiveSeq triggered by onClick on the video shape: cmd="togglePause" (presetID 2)
               python-pptx's _add_video_timing writes only the p:video node, so the click does nothing.

The poster image is what PDF exports, LibreOffice and offline viewers show; playback needs PowerPoint + internet.

Usage
    from pptx_online_video import add_online_video
    add_online_video(slide, "https://www.youtube.com/shorts/gdHwOn2kvT8", x=10.5, y=2.2, w=2.2, h=3.9,
                     poster=None, name="How to set up your first OpenAI Dot")  # poster=None -> YouTube thumbnail

Upgrade an existing deck whose videos only show a poster:
    python3 pptx_online_video.py fix deck.pptx [out.pptx]
"""
import io, re, sys, urllib.request

from lxml import etree
from pptx.util import Inches
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
PNS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
MEDIA_RT = 'http://schemas.microsoft.com/office/2007/relationships/media'


def youtube_id(url):
    m = re.search(r'(?:youtu\.be/|/shorts/|/embed/|[?&]v=)([\w-]{11})', url)
    if not m:
        raise ValueError(f'not a YouTube URL: {url}')
    return m.group(1)


def embed_url(url):
    return f'https://www.youtube.com/embed/{youtube_id(url)}?feature=oembed'


def thumbnail(url):
    """YouTube's own poster frame (bytes); Shorts use the portrait oar2.jpg when available."""
    vid = youtube_id(url)
    for name in (('oar2.jpg', 'hqdefault.jpg') if '/shorts/' in url else ('maxresdefault.jpg', 'hqdefault.jpg')):
        try:
            with urllib.request.urlopen(f'https://i.ytimg.com/vi/{vid}/{name}', timeout=20) as r:
                return r.read()
        except Exception:
            continue
    raise RuntimeError(f'no thumbnail for {url}; pass poster=')


def _video_spids(slide):
    """Shape ids of every online video on the slide, in z-order."""
    ids = []
    for vf in slide.shapes._spTree.iter(qn('a:videoFile')):
        pic = vf.getparent().getparent().getparent()
        ids.append(pic.find(f'{P}nvPicPr/{P}cNvPr').get('id'))
    return ids


def _timing_xml(spids):
    """PowerPoint's click-to-play timing for the given video shape ids (one main-sequence play + one
    click trigger per video), exactly as PowerPoint for Mac writes it for an Online Film."""
    n = [2]  # cTn ids: tmRoot=1, mainSeq=2, then sequential
    def nid():
        n[0] += 1
        return n[0]
    seq_id = 2
    main = []
    for sp in spids:
        a, b, c, d = nid(), nid(), nid(), nid()
        main.append(
            f'<p:par><p:cTn id="{a}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{b}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{c}" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="clickEffect">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr>'
            f'<p:cTn id="{d}" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{sp}"/></p:tgtEl></p:cBhvr></p:cmd>'
            f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    tail = []
    for sp in spids:
        tail.append(f'<p:video><p:cMediaNode vol="80000"><p:cTn id="{nid()}" fill="hold" display="0"><p:stCondLst>'
                    f'<p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{sp}"/></p:tgtEl>'
                    f'</p:cMediaNode></p:video>')
    for sp in spids:
        s, a, b, c, d = nid(), nid(), nid(), nid(), nid()
        tail.append(
            f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="{s}" restart="whenNotActive" fill="hold" evtFilter="cancelBubble" '
            f'nodeType="interactiveSeq"><p:stCondLst><p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="{sp}"/></p:tgtEl>'
            f'</p:cond></p:stCondLst><p:endSync evt="end" delay="0"><p:rtn val="all"/></p:endSync><p:childTnLst>'
            f'<p:par><p:cTn id="{a}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{b}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{c}" presetID="2" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="clickEffect">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:cmd type="call" cmd="togglePause"><p:cBhvr>'
            f'<p:cTn id="{d}" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{sp}"/></p:tgtEl></p:cBhvr></p:cmd>'
            f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst>'
            f'</p:cTn><p:nextCondLst><p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="{sp}"/></p:tgtEl></p:cond>'
            f'</p:nextCondLst></p:seq>')
    return (f'<p:timing xmlns:p="{PNS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="{seq_id}" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
            + ''.join(main) +
            '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
            '</p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
            '</p:nextCondLst></p:seq>' + ''.join(tail) + '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')


def set_video_timing(slide):
    """(Re)write the slide's p:timing so every online video on it plays on click. Replaces existing timing,
    so call it after the slide's videos are added (add_online_video does this for you)."""
    sld = slide._element
    for t in sld.findall(f'{P}timing'):
        sld.remove(t)
    spids = _video_spids(slide)
    if not spids:
        return
    timing = etree.fromstring(_timing_xml(spids))
    # schema order: cSld, clrMapOvr, transition, timing, extLst
    anchor = sld.find(f'{P}extLst')
    if anchor is not None:
        anchor.addprevious(timing)
    else:
        sld.append(timing)


def _clean_nvPr(nvPr):
    for e in nvPr.findall(f'{P}extLst'):   # PowerPoint's Online Film has no p14:media / p15:webVideoPr here
        nvPr.remove(e)


def add_online_video(slide, url, x, y, w, h, poster=None, name='Online video'):
    """Add a click-to-play YouTube video (inches). poster: path or file-like; None downloads the thumbnail."""
    from pptx.oxml.shapes.picture import CT_Picture
    embed = embed_url(url)
    part = slide.part
    v_rId = part.relate_to(embed, RT.VIDEO, is_external=True)
    _, p_rId = part.get_or_add_image_part(poster if poster is not None else io.BytesIO(thumbnail(url)))
    pic = CT_Picture.new_video_pic(slide.shapes._next_shape_id, name, v_rId, 'rIdUnused', p_rId,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    cNvPr = pic.find(f'{P}nvPicPr/{P}cNvPr')
    cNvPr.set('descr', name)
    locks = pic.find(f'{P}nvPicPr/{P}cNvPicPr').find(qn('a:picLocks'))
    if locks is not None:
        locks.set('noRot', '1')
    _clean_nvPr(pic.find(f'{P}nvPicPr/{P}nvPr'))
    slide.shapes._spTree.append(pic)
    set_video_timing(slide)
    return pic


def fix_deck(src, dst=None):
    """Upgrade every YouTube video in a deck to PowerPoint's click-to-play Online Film form; returns count fixed."""
    from pptx import Presentation
    prs = Presentation(src)
    n = 0
    for slide in prs.slides:
        part = slide.part
        found = 0
        for vf in list(slide.shapes._spTree.iter(qn('a:videoFile'))):
            rId = vf.get(qn('r:link'))
            if rId not in part.rels or 'youtu' not in part.rels[rId].target_ref:
                continue
            _clean_nvPr(vf.getparent())
            found += 1
        if found:
            for rId, rel in list(part.rels.items()):      # drop now-unused linked-media relationships
                if rel.reltype == MEDIA_RT and rel.is_external and 'youtu' in rel.target_ref:
                    part.drop_rel(rId)
            set_video_timing(slide)
            n += found
    prs.save(dst or src)
    return n


if __name__ == '__main__':
    if len(sys.argv) >= 3 and sys.argv[1] == 'fix':
        print(f'{fix_deck(sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)} YouTube video(s) made playable')
    else:
        print(__doc__)
