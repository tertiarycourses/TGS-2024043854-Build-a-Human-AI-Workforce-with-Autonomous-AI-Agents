"""
ccdeck — slide components in the "Agentic AI Applications with Claude Code"
(TGS-2025052468) house design: warm ivory (FAF9F5) and charcoal (21201C)
canvases, terracotta accent (D97757), Cambria headings, Calibri body, white
rounded cards with a soft shadow and an icon in a filled circle.

Every component draws with native PowerPoint shapes (editable by the trainer)
and keeps text inside its box via a conservative fit-to-box font sizing.
Icons are react-icons glyphs rendered to PNG by render_icons.js and cached in
./icons/.
"""
import math, os, subprocess
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ICON_DIR = os.path.join(HERE, "icons")

# ---------------------------------------------------------------- palette
IVORY = "FAF9F5"; CHAR = "21201C"; PANEL = "2E2C26"; WHITE = "FFFFFF"
ACC = "D97757"; ACC_D = "C2603F"; TEAL = "3A7D7B"; PLUM = "6D2E46"; OLIVE = "7A8B4F"
INK = "21201C"; MUTED = "6B6862"; SOFT = "C9C4B9"; LIGHT = "E8E5DE"; DIM = "B9B4A9"
FOOT = "8A867D"; RULE = "E3DFD6"; CODEBG = "2B2924"
SERIF = "Cambria"; SANS = "Calibri"; MONO = "Consolas"

def rgb(h): return RGBColor.from_string(h)

# ---------------------------------------------------------------- icons
ICONS = {
 # name: react-icons pack/component
 "robot":"fa6/FaRobot","brain":"lu/LuBrain","eye":"lu/LuEye","hand":"lu/LuHand","bolt":"lu/LuZap",
 "users":"lu/LuUsers","user":"lu/LuUser","handshake":"lu/LuHandshake","globe":"lu/LuGlobe",
 "memory":"lu/LuDatabase","book":"lu/LuBookOpen","search":"lu/LuSearch","target":"lu/LuTarget",
 "loop":"lu/LuRefreshCw","check":"lu/LuCircleCheck","shield":"lu/LuShieldCheck","lock":"lu/LuLock",
 "tool":"lu/LuWrench","plug":"lu/LuPlug","clock":"lu/LuClock","calendar":"lu/LuCalendarDays",
 "chat":"lu/LuMessageSquare","send":"lu/LuSend","video":"lu/LuClapperboard","film":"lu/LuFilm",
 "pen":"lu/LuPenLine","image":"lu/LuImage","chart":"lu/LuChartColumn","trend":"lu/LuTrendingUp",
 "flask":"lu/LuFlaskConical","pill":"lu/LuPill","sigma":"lu/LuSigma","atom":"lu/LuAtom",
 "heart":"lu/LuHeartPulse","bank":"lu/LuLandmark","code":"lu/LuCode","cart":"lu/LuShoppingCart",
 "megaphone":"lu/LuMegaphone","building":"lu/LuBuilding2","org":"lu/LuNetwork","layers":"lu/LuLayers",
 "list":"lu/LuListChecks","kanban":"lu/LuSquareKanban","coins":"lu/LuCoins","alert":"lu/LuTriangleAlert",
 "flag":"lu/LuFlag","rocket":"lu/LuRocket","map":"lu/LuMap","cube":"lu/LuBox","terminal":"lu/LuTerminal",
 "download":"lu/LuDownload","login":"lu/LuLogIn","phone":"lu/LuSmartphone","clipboard":"lu/LuClipboardCheck",
 "gavel":"lu/LuGavel","scale":"lu/LuScale","sparkles":"lu/LuSparkles","lightbulb":"lu/LuLightbulb",
 "puzzle":"lu/LuPuzzle","gear":"lu/LuSettings","history":"lu/LuHistory","file":"lu/LuFileText",
 "folder":"lu/LuFolderOpen","mail":"lu/LuMail","star":"lu/LuStar","thumbs":"lu/LuThumbsUp",
 "pause":"lu/LuCirclePause","eyecheck":"lu/LuScanEye","route":"lu/LuRoute","gauge":"lu/LuGauge",
 "medal":"lu/LuAward","linkedin":"fa6/FaLinkedin","github":"fa6/FaGithub","play":"lu/LuCirclePlay",
 "grad":"lu/LuGraduationCap","briefcase":"lu/LuBriefcase","question":"lu/LuCircleHelp",
 "chess":"fa6/FaChessKnight","network":"lu/LuWaypoints","dna":"lu/LuDna","hospital":"lu/LuHospital",
 "camera":"lu/LuCamera","hash":"lu/LuHash","inbox":"lu/LuInbox","edit":"lu/LuFilePen",
 "coffee":"lu/LuCoffee","store":"lu/LuStore","wallet":"lu/LuWallet","smile":"lu/LuSmile",
 "server":"lu/LuServer","box":"lu/LuPackage","stop":"lu/LuOctagonX","cpu":"lu/LuCpu","wand":"lu/LuWandSparkles",
}

def icon_png(name, hexc=WHITE):
    os.makedirs(ICON_DIR, exist_ok=True)
    path = os.path.join(ICON_DIR, f"{name}-{hexc}.png")
    if not os.path.exists(path):
        env = dict(os.environ)
        root = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
        env["NODE_PATH"] = root
        subprocess.run(["node", os.path.join(HERE, "render_icons.js"), ICON_DIR,
                        f"{name}:{ICONS[name]}:{hexc}"], check=True, env=env)
    return path

# ---------------------------------------------------------------- deck
class Deck:
    def __init__(self, title, code, short):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333); self.prs.slide_height = Inches(7.5)
        self.blank = self.prs.slide_layouts[6]
        self.title = title; self.code = code; self.short = short
        self.dark = False
        self.n = 0

    # ---------- primitives
    def new(self, dark=False):
        s = self.prs.slides.add_slide(self.blank)
        self.dark = dark
        s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(CHAR if dark else IVORY)
        self.n += 1
        return s

    def save(self, path):
        trans = ('<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                 'spd="med"><p:fade/></p:transition>')
        for s in self.prs.slides:
            if s.element.find(qn('p:transition')) is None:
                s.element.append(etree.fromstring(trans))
        self.prs.save(path)

# ---------------------------------------------------------------- shape helpers
def _shadow(shape):
    spPr = shape._element.spPr
    eff = etree.SubElement(spPr, qn('a:effectLst'))
    sh = etree.SubElement(eff, qn('a:outerShdw'), blurRad="101600", dist="38100", dir="5400000",
                          algn="bl", rotWithShape="0")
    c = etree.SubElement(sh, qn('a:srgbClr'), val="000000")
    etree.SubElement(c, qn('a:alpha'), val="16000")

def box(s, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, radius=None, shadow=False):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
    else: sp.fill.background()
    if line: sp.line.color.rgb = rgb(line); sp.line.width = Pt(lw)
    else: sp.line.fill.background()
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius
    if shadow: _shadow(sp)
    sp.text_frame.text = ""
    return sp

def card(s, x, y, w, h, fill=WHITE, shadow=True, line=None):
    return box(s, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.035, shadow=shadow)

def circle(s, x, y, d, fill, line=None):
    return box(s, x, y, d, d, fill=fill, line=line, shape=MSO_SHAPE.OVAL)

def line(s, x1, y1, x2, y2, color=MUTED, w=1.5, arrow=False, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn('a:prstDash'), val="dash")
    if arrow:
        etree.SubElement(ln, qn('a:tailEnd'), type="triangle", w="med", len="med")
    return c

def fit(text, w, h, maxpt, minpt=8, lh=1.18, cw=0.5):
    """Largest point size at which `text` fits a w x h inch box (estimate)."""
    paras = str(text).split("\n")
    pt = maxpt
    while pt > minpt:
        cpl = max(1, int((w - 0.1) * 72 / (pt * cw)))
        lines = sum(max(1, math.ceil(len(p) / cpl)) for p in paras)
        if lines * pt * lh / 72 <= h - 0.06: return pt
        pt -= 0.5
    return minpt

ARROWS = "→←↔⇄↑↓"

def _runs(p, t, font, size, bold, italic, color):
    """Add runs to paragraph p; arrow glyphs use Calibri (Cambria/Caladea lack them)."""
    buf = ""; cur = None
    def flush(b, f):
        if not b: return
        r = p.add_run(); r.text = b; fo = r.font
        fo.size = Pt(size); fo.bold = bold; fo.italic = italic; fo.name = f; fo.color.rgb = rgb(color)
    for ch in t:
        f = SANS if (ch in ARROWS and font == SERIF) else font
        if cur is not None and f != cur:
            flush(buf, cur); buf = ""
        cur = f; buf += ch
    flush(buf, cur or font)

def text(s, x, y, w, h, t, size=14, color=INK, bold=False, font=SANS, align="l",
         anchor="t", italic=False, autofit=True, minpt=8, lh=1.12, cw=0.5):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if autofit: size = fit(t, w, h, size, minpt, cw=cw)
    for i, para in enumerate(str(t).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = lh
        _runs(p, para, font, size, bold, italic, color)
    return tb

def rich(s, x, y, w, h, runs_paras, size=14, align="l", anchor="t", lh=1.15, space=4):
    """runs_paras: list of paragraphs, each a list of (text, dict(color,bold,font,size,italic))."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, runs in enumerate(runs_paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = lh; p.space_after = Pt(space)
        for t, st in runs:
            _runs(p, t, st.get("font", SANS), st.get("size", size), st.get("bold", False),
                  st.get("italic", False), st.get("color", INK))
    return tb

def icon(s, name, x, y, d, fill=ACC, glyph=WHITE):
    circle(s, x, y, d, fill)
    g = d * 0.52
    s.shapes.add_picture(icon_png(name, glyph), Inches(x + (d - g) / 2), Inches(y + (d - g) / 2), Inches(g), Inches(g))

def picture(s, path, x, y, w=None, h=None, frame=True, maxw=None, maxh=None):
    """Place an image; if maxw/maxh given, fit inside that box (centred)."""
    from PIL import Image
    iw, ih = Image.open(path).size
    if maxw and maxh:
        sc = min(maxw / iw, maxh / ih); w, h = iw * sc, ih * sc
        x = x + (maxw - w) / 2; y = y + (maxh - h) / 2
    if frame:
        card(s, x - 0.06, y - 0.06, w + 0.12, h + 0.12, fill=WHITE)
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    return x, y, w, h

# ---------------------------------------------------------------- chrome
def footer(D, s):
    c = FOOT if D.dark else MUTED
    # page number first: text extractors that join shapes without spaces must not glue it onto the code
    text(s, 12.4, 7.08, 0.55, 0.3, str(D.n), 9, c, align="r", autofit=False)
    text(s, 0.5, 7.08, 7.0, 0.3, "© Tertiary Infotech Academy Pte Ltd", 9, c, autofit=False)
    text(s, 6.4, 7.08, 5.9, 0.3, f"{D.short}  ·  {D.code}", 9, c, align="r", autofit=False)

def header(D, s, kicker, title, sub=None):
    dark = D.dark
    text(s, 0.6, 0.42, 11.5, 0.3, kicker.upper(), 12, ACC, bold=True, autofit=False)
    long = len(title) > 58
    if long:   # two-line title: smaller type, subtitle pushed down so they never crowd
        text(s, 0.6, 0.68, 12.1, 1.02, title, 24, WHITE if dark else INK, bold=True, font=SERIF, cw=0.52, lh=1.0)
    else:
        text(s, 0.6, 0.72, 12.1, 0.8, title, 30, WHITE if dark else INK, bold=True, font=SERIF, cw=0.52)
    if sub:
        text(s, 0.6, 1.76 if long else 1.62, 12.1, 0.6 if long else 0.62, sub, 15 if long else 16, SOFT if dark else INK, minpt=11)
    footer(D, s)

def tip(D, s, t, y=6.4, label="TIP"):
    rich(s, 0.6, y, 12.1, 0.5, [[(label + "  ", dict(color=ACC, bold=True, size=13)),
                                 (t, dict(color=SOFT if D.dark else MUTED, size=13))]])

# ---------------------------------------------------------------- slide types
def cover(D, kicker, title, version_line, org_lines, contact, logo=None):
    s = D.new(dark=True)
    circle(s, 10.4, -1.6, 4.6, ACC); circle(s, 11.6, 4.4, 3.4, TEAL)
    if logo and os.path.exists(logo):
        from PIL import Image
        iw, ih = Image.open(logo).size; h = 0.62; w = h * iw / ih
        card(s, 0.82, 0.42, w + 0.3, h + 0.2, fill=WHITE, shadow=False)
        s.shapes.add_picture(logo, Inches(0.97), Inches(0.52), Inches(w), Inches(h))
    text(s, 0.5, 7.08, 7.0, 0.3, "© Tertiary Infotech Academy Pte Ltd. All rights reserved.", 9, FOOT, autofit=False)
    text(s, 0.9, 1.5, 8.5, 0.4, kicker, 14, ACC, bold=True, autofit=False)
    text(s, 0.85, 2.0, 10.4, 2.6, title, 48, WHITE, bold=True, font=SERIF, cw=0.52, lh=1.02)
    text(s, 0.9, 4.75, 8.0, 0.4, version_line, 14, ACC, bold=True, autofit=False)
    text(s, 0.9, 5.55, 9.5, 0.8, org_lines, 15, WHITE, bold=True, autofit=False)
    text(s, 0.9, 6.7, 11.5, 0.35, contact, 12, DIM, autofit=False)
    return s

def section(D, label, title, items, sub=None):
    s = D.new(dark=True)
    circle(s, -1.7, 5.0, 4.4, ACC)
    text(s, 0.9, 2.0, 7.6, 0.5, label.upper(), 18, ACC, bold=True, autofit=False)
    text(s, 0.85, 2.55, 7.9, 2.2, title, 40, WHITE, bold=True, font=SERIF, cw=0.52, lh=1.04)
    if sub: text(s, 0.9, 4.75, 7.6, 0.9, sub, 15, SOFT)
    n = len(items); rowh = 0.56 if n <= 6 else 0.46
    ph = 0.45 + n * rowh
    py = max(0.6, (7.5 - ph) / 2 - 0.2)
    card(s, 8.95, py, 3.85, ph, fill=PANEL, shadow=False)
    for i, it in enumerate(items):
        yy = py + 0.25 + i * rowh
        text(s, 9.15, yy, 0.5, 0.42, str(i + 1), 17, ACC, bold=True, font=SERIF, autofit=False)
        text(s, 9.65, yy, 3.0, rowh, it, 13.5, LIGHT, anchor="t", minpt=10)
    text(s, 6.4, 7.08, 6.4, 0.3, f"© Tertiary Infotech Academy Pte Ltd  ·  {D.code}", 9, FOOT, align="r", autofit=False)
    return s

def cards(D, kicker, title, items, sub=None, cols=3, dark=False, tipt=None, icons=None,
          hi=(), top=None, bottom=6.3, accent=ACC, desc_pt=14):
    """items: list of (title, desc) or (title, desc, icon). hi: indices drawn as dark cards."""
    s = D.new(dark=dark); header(D, s, kicker, title, sub)
    n = len(items); rows = math.ceil(n / cols)
    top = top or (2.45 if sub else 1.85)
    gap = 0.2; W = 12.1; cw = (W - gap * (cols - 1)) / cols
    ch = (bottom - top - gap * (rows - 1)) / rows
    for i, it in enumerate(items):
        t, d = it[0], it[1]; ic = it[2] if len(it) > 2 else (icons[i] if icons else None)
        r, c = divmod(i, cols)
        x = 0.6 + c * (cw + gap); y = top + r * (ch + gap)
        darkc = (i in hi) != dark
        card(s, x, y, cw, ch, fill=(CHAR if (i in hi and not dark) else PANEL) if darkc else WHITE,
             shadow=not dark)
        pad = 0.25; yy = y + 0.2
        if ic:
            dd = 0.5 if ch < 1.8 else 0.56
            icon(s, ic, x + pad, yy, dd, fill=accent if not darkc else ACC)
            if ch < 2.15:  # compact: title beside icon
                text(s, x + pad + dd + 0.15, yy - 0.02, cw - 2 * pad - dd - 0.15, dd + 0.04, t, 14,
                     WHITE if darkc else INK, bold=True, font=SERIF, anchor="m", minpt=10)
                text(s, x + pad, yy + dd + 0.08, cw - 2 * pad, y + ch - (yy + dd + 0.08) - 0.08, d,
                     desc_pt, SOFT if darkc else MUTED, minpt=8.5)
                continue
            yy += dd + 0.14
        th = 0.42 if len(t) < 30 or cw > 4 else 0.62
        text(s, x + pad, yy, cw - 2 * pad, th, t, 15, WHITE if darkc else INK, bold=True, font=SERIF, minpt=10.5)
        yy += th + 0.04
        text(s, x + pad, yy, cw - 2 * pad, y + ch - yy - 0.12, d, desc_pt, SOFT if darkc else MUTED, minpt=8.5)
    if tipt: tip(D, s, tipt, y=bottom + 0.12)
    return s

def rows(D, kicker, title, items, sub=None, dark=True, tipt=None, accent=ACC, top=None):
    """Vertical list: icon circle + bold title + description. items: (icon,title,desc)."""
    s = D.new(dark=dark); header(D, s, kicker, title, sub)
    top = top or (2.4 if sub else 1.9)
    n = len(items); avail = (6.25 if tipt else 6.85) - top
    rh = min(1.08, avail / n)
    for i, (ic, t, d) in enumerate(items):
        y = top + i * rh
        icon(s, ic, 0.7, y + 0.04, min(0.6, rh - 0.12), fill=accent)
        text(s, 1.5, y - 0.02, 11.2, 0.42, t, 18 if rh > 0.85 else 16, WHITE if dark else INK, bold=True, font=SERIF, minpt=12)
        text(s, 1.5, y + 0.38, 11.2, rh - 0.42, d, 14 if rh > 0.85 else 12.5, SOFT if dark else MUTED, minpt=9.5)
    if tipt: tip(D, s, tipt, y=6.35)
    return s

def timeline(D, kicker, title, nodes, sub=None, foot=None, dark=True):
    """nodes: (label, year, desc). Horizontal timeline with year circles."""
    s = D.new(dark=dark); header(D, s, kicker, title, sub)
    n = len(nodes); x0, x1 = 0.6, 12.7; w = (x1 - x0) / n
    ly = 3.75 if not sub else 4.05
    line(s, x0 + w / 2, ly, x1 - w / 2, ly, MUTED if dark else SOFT, 2)
    d = 0.86 if n <= 5 else (0.74 if n <= 7 else 0.64)
    for i, (lab, yr, desc) in enumerate(nodes):
        cx = x0 + w * i + w / 2
        text(s, cx - w / 2 + 0.03, ly - 1.45, w - 0.06, 0.95, lab, 15 if n <= 5 else 12.5, ACC, bold=True,
             font=SERIF, align="c", anchor="b", minpt=9)
        circle(s, cx - d / 2, ly - d / 2, d, ACC)
        text(s, cx - d / 2, ly - d / 2, d, d, yr, 14 if n <= 6 else 11.5, WHITE, bold=True, font=SERIF,
             align="c", anchor="m", autofit=False)
        text(s, cx - w / 2 + 0.03, ly + d / 2 + 0.15, w - 0.06, 1.75, desc, 12 if n <= 6 else 10.5,
             SOFT if dark else MUTED, align="c", minpt=8.5)
    if foot: text(s, 0.6, 6.45, 12.1, 0.42, foot, 13, ACC, align="c", minpt=10)
    return s

def loop(D, kicker, title, steps, sub=None, foot=None, dark=False):
    """steps: (icon,title,desc). Alternating white/dark cards with arrows — CC agentic-loop look."""
    s = D.new(dark=dark); header(D, s, kicker, title, sub)
    n = len(steps); gap = 0.32; W = 12.1; cw = (W - gap * (n - 1)) / n
    top = 2.6 if sub else 2.1; ch = 3.1
    for i, (ic, t, d) in enumerate(steps):
        x = 0.6 + i * (cw + gap); darkc = (i % 2 == 1) != dark
        card(s, x, top, cw, ch, fill=(CHAR if not dark else PANEL) if darkc else WHITE, shadow=not dark)
        icon(s, ic, x + 0.28, top + 0.28, 0.6, fill=ACC if darkc else TEAL)
        text(s, x + 0.28, top + 1.0, cw - 0.5, 0.62, t, 16, WHITE if darkc else INK, bold=True, font=SERIF, minpt=11)
        text(s, x + 0.28, top + 1.62, cw - 0.5, ch - 1.72, d, 12.5, SOFT if darkc else MUTED, minpt=9)
        if i < n - 1:
            line(s, x + cw + 0.04, top + ch / 2, x + cw + gap - 0.04, top + ch / 2, ACC, 2, arrow=True)
    if foot: text(s, 0.6, top + ch + 0.3, 12.1, 0.42, foot, 14, ACC_D, align="c", minpt=10)
    return s

def compare(D, kicker, title, left, right, sub=None, tipt=None):
    """left/right: (heading, icon, [bullets]). Left white card, right dark card."""
    s = D.new(); header(D, s, kicker, title, sub)
    top = 2.45 if sub else 1.95; bot = 6.15 if tipt else 6.75
    for j, (hd, ic, bl) in enumerate([left, right]):
        x = 0.6 + j * 6.15; darkc = j == 1
        card(s, x, top, 5.95, bot - top, fill=CHAR if darkc else WHITE)
        icon(s, ic, x + 0.3, top + 0.3, 0.6, fill=ACC)
        text(s, x + 1.05, top + 0.32, 4.7, 0.56, hd, 19, WHITE if darkc else INK, bold=True, font=SERIF, anchor="m", minpt=12)
        body = "\n".join("•  " + b for b in bl)
        text(s, x + 0.35, top + 1.15, 5.3, bot - top - 1.3, body, 16.5, SOFT if darkc else MUTED, minpt=9.5, lh=1.2)
    if tipt: tip(D, s, tipt, y=6.3)
    return s

def table(D, kicker, title, head, body, sub=None, colw=None, tipt=None, size=11.5, dark=False, hi_col=None):
    s = D.new(dark=dark); header(D, s, kicker, title, sub)
    top = 2.4 if sub else 1.85; bot = 6.2 if tipt else 6.75
    nrows = len(body) + 1; ncols = len(head)
    colw = colw or [12.1 / ncols] * ncols
    tot = sum(colw); colw = [c * 12.1 / tot for c in colw]
    gt = s.shapes.add_table(nrows, ncols, Inches(0.6), Inches(top), Inches(12.1), Inches(bot - top)).table
    tblPr = gt._tbl.tblPr
    for k in ("bandRow", "firstRow"): tblPr.set(k, "0")
    for j, w in enumerate(colw): gt.columns[j].width = Inches(w)
    rh = (bot - top) / nrows
    # grow the body font while every cell still fits its row
    def _fits(pt):
        for row in body:
            for j, v in enumerate(row):
                cpl = max(1, int((colw[j] - 0.16) * 72 / (pt * 0.5)))
                if math.ceil(len(str(v)) / cpl) * pt * 1.2 / 72 > rh - 0.08: return False
        return True
    while size < 15 and _fits(size + 0.5): size += 0.5
    for i in range(nrows):
        gt.rows[i].height = Inches(rh if i else min(rh, 0.5))
        for j in range(ncols):
            c = gt.cell(i, j); val = head[j] if i == 0 else body[i - 1][j]
            c.fill.solid()
            if i == 0: c.fill.fore_color.rgb = rgb(CHAR if not dark else ACC_D)
            elif hi_col is not None and j == hi_col: c.fill.fore_color.rgb = rgb("F6E7DF")
            else: c.fill.fore_color.rgb = rgb(WHITE if i % 2 else "F3F1EB")
            c.margin_left = c.margin_right = Inches(0.07); c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True; tf.text = ""
            p = tf.paragraphs[0]; r = p.add_run(); r.text = str(val)
            f = r.font; f.name = SERIF if (i == 0 or j == 0) else SANS
            f.bold = (i == 0 or j == 0); f.size = Pt(size + (0.5 if i == 0 else 0))
            f.color.rgb = rgb(WHITE if i == 0 else (INK if j == 0 else MUTED))
    if tipt: tip(D, s, tipt, y=6.3)
    return s

def image(D, kicker, title, path, caption=None, sub=None, side=None, dark=False, src=None):
    """Screenshot slide. side: list of (icon,title,desc) shown right of the image."""
    s = D.new(dark=dark); header(D, s, kicker, title, sub)
    top = 2.4 if sub else 1.85; bot = 6.35 if caption else 6.85
    if side:
        picture(s, path, 0.7, top + 0.05, maxw=7.3, maxh=bot - top - 0.1)
        n = len(side); rh = (bot - top) / n
        for i, (ic, t, d) in enumerate(side):
            y = top + i * rh
            icon(s, ic, 8.35, y + 0.05, 0.5)
            text(s, 9.0, y, 3.75, 0.4, t, 15, WHITE if dark else INK, bold=True, font=SERIF, minpt=11)
            text(s, 9.0, y + 0.4, 3.75, rh - 0.45, d, 12, SOFT if dark else MUTED, minpt=9)
    else:
        picture(s, path, 0.7, top + 0.05, maxw=11.9, maxh=bot - top - 0.1)
    if caption:
        rich(s, 0.6, 6.42, 12.1, 0.55, [[("▸  ", dict(color=ACC, bold=True, size=13)),
                                        (caption, dict(color=SOFT if dark else MUTED, size=12.5))]])
    if src: text(s, 0.6, 6.8, 12.1, 0.25, src, 9, FOOT, autofit=False)
    return s

def statement(D, kicker, big, small=None, dark=True):
    s = D.new(dark=dark)
    circle(s, 10.6, -1.2, 3.8, ACC)
    text(s, 0.9, 1.9, 9.5, 0.4, kicker.upper(), 14, ACC, bold=True, autofit=False)
    text(s, 0.85, 2.4, 10.0, 2.2, big, 40, WHITE if dark else INK, bold=True, font=SERIF, cw=0.52, lh=1.05)
    if small: text(s, 0.9, 4.7, 10.0, 1.5, small, 17, SOFT if dark else MUTED)
    footer(D, s)
    return s

def prompt(D, kicker, title, prompt_text, sub=None, side=None, label="PROMPT  ·  copy from the Lab Prompts page", dark=False):
    """A copy-ready prompt panel (dark code card) + optional right-side notes."""
    s = D.new(dark=dark); header(D, s, kicker, title, sub)
    top = 2.4 if sub else 1.9; bot = 6.8
    pw = 7.7 if side else 12.1
    card(s, 0.6, top, pw, bot - top, fill=CODEBG, shadow=not dark)
    for k, c in enumerate(("E06C5A", "E5B454", "7DB36B")):
        circle(s, 0.82 + k * 0.24, top + 0.18, 0.13, c)
    text(s, 1.6, top + 0.1, pw - 1.2, 0.3, label, 10, DIM, autofit=False)
    text(s, 0.85, top + 0.5, pw - 0.5, bot - top - 0.62, prompt_text, 13, "F1EDE4", font=MONO, minpt=8, cw=0.66, lh=1.1)
    if side:
        n = len(side); rh = (bot - top) / n
        for i, (ic, t, d) in enumerate(side):
            y = top + i * rh
            icon(s, ic, 8.6, y + 0.04, 0.5)
            text(s, 9.25, y, 3.5, 0.4, t, 15, WHITE if dark else INK, bold=True, font=SERIF, minpt=11)
            text(s, 9.25, y + 0.4, 3.5, rh - 0.45, d, 12, SOFT if dark else MUTED, minpt=9)
    return s

def lab_intro(D, lab, kicker):
    """Lab overview: scenario on the left (dark), deliverables + agents + time on the right."""
    s = D.new(); header(D, s, kicker, f"Lab {lab['num']} — {lab['title']}", lab["hook"])
    top = 2.45
    card(s, 0.6, top, 6.3, 4.3, fill=CHAR)
    icon(s, lab.get("icon", "rocket"), 0.9, top + 0.3, 0.6)
    text(s, 1.65, top + 0.32, 5.0, 0.56, "The business scenario", 18, WHITE, bold=True, font=SERIF, anchor="m")
    text(s, 0.9, top + 1.1, 5.75, 3.05, lab["scenario"], 14, SOFT, minpt=9.5, lh=1.2)
    # right column
    card(s, 7.1, top, 5.6, 2.55, fill=WHITE)
    text(s, 7.35, top + 0.18, 5.1, 0.4, "You will deliver", 16, INK, bold=True, font=SERIF)
    text(s, 7.35, top + 0.62, 5.15, 1.85, "\n".join("✓  " + d for d in lab["deliverables"]), 12.5, MUTED, minpt=9, lh=1.15)
    card(s, 7.1, top + 2.72, 5.6, 1.58, fill=WHITE)
    rich(s, 7.35, top + 2.86, 5.15, 1.35, [
        [("Agents  ", dict(color=ACC, bold=True, size=12.5)), (lab["agents"], dict(color=MUTED, size=12.5))],
        [("Human role  ", dict(color=ACC, bold=True, size=12.5)), (lab["human"], dict(color=MUTED, size=12.5))],
        [("Time  ", dict(color=ACC, bold=True, size=12.5)), (f"{lab['minutes']} min  ·  {lab['paradigm']}  ·  labs/{lab['slug']}/", dict(color=MUTED, size=12.5))],
    ], size=12.5, space=3)
    return s

def lab_flow(D, lab, kicker):
    """The lab's step flow — numbered chevron-like cards (max 6)."""
    steps = lab["flow"]
    s = D.new(dark=True); header(D, s, kicker, f"Lab {lab['num']} — How the workflow runs",
                                 "Perception → reasoning → action, with you at the checkpoints.")
    n = len(steps); gap = 0.18; W = 12.1; cw = (W - gap * (n - 1)) / n; top = 2.55; ch = 3.4
    for i, (ic, t, d, who) in enumerate(steps):
        x = 0.6 + i * (cw + gap); human = who == "human"
        card(s, x, top, cw, ch, fill=ACC_D if human else PANEL, shadow=False)
        text(s, x + 0.2, top + 0.15, cw - 0.4, 0.3, ("HUMAN" if human else "AGENT") + f"  ·  STEP {i+1}", 9.5,
             WHITE if human else ACC, bold=True, autofit=False)
        icon(s, ic, x + 0.2, top + 0.55, 0.52, fill=WHITE if human else ACC, glyph=ACC_D if human else WHITE)
        text(s, x + 0.2, top + 1.18, cw - 0.35, 0.62, t, 14.5, WHITE, bold=True, font=SERIF, minpt=10)
        text(s, x + 0.2, top + 1.8, cw - 0.35, ch - 1.9, d, 13.5, LIGHT if human else SOFT, minpt=8.5)
        if i < n - 1:
            line(s, x + cw + 0.01, top + ch / 2, x + cw + gap - 0.01, top + ch / 2, ACC, 1.75, arrow=True)
    rich(s, 0.6, 6.25, 12.1, 0.5, [[("Checkpoint  ", dict(color=ACC, bold=True, size=13)),
                                   (lab["check"], dict(color=SOFT, size=13))]])
    return s

def hub(D, kicker, title, center, spokes, sub=None, foot=None, dark=False):
    """Hub-and-spoke diagram: centre node (orchestrator) + up to 6 spokes (icon,title,desc)."""
    s = D.new(dark=dark); header(D, s, kicker, title, sub)
    cx, cy = 6.67, 4.3 if sub else 4.1
    n = len(spokes); R_x, R_y = 4.35, 1.75
    pos = []
    for i in range(n):
        a = -math.pi / 2 + 2 * math.pi * i / n
        pos.append((cx + R_x * math.cos(a), cy + R_y * math.sin(a)))
    for (px, py) in pos:
        line(s, cx, cy, px, py, SOFT, 1.5, dash=True)
    cw, chh = 2.9, 1.0
    card(s, cx - cw / 2, cy - chh / 2, cw, chh, fill=ACC, shadow=True)
    text(s, cx - cw / 2, cy - chh / 2, cw, chh, center, 16, WHITE, bold=True, font=SERIF, align="c", anchor="m", minpt=11)
    for (px, py), (ic, t, d) in zip(pos, spokes):
        w, h = 2.75, 1.05
        card(s, px - w / 2, py - h / 2, w, h, fill=PANEL if dark else WHITE)
        icon(s, ic, px - w / 2 + 0.14, py - h / 2 + 0.14, 0.42, fill=TEAL)
        text(s, px - w / 2 + 0.64, py - h / 2 + 0.1, w - 0.74, 0.36, t, 13, WHITE if dark else INK, bold=True, font=SERIF, minpt=9.5)
        text(s, px - w / 2 + 0.14, py - h / 2 + 0.58, w - 0.26, 0.45, d, 10, SOFT if dark else MUTED, minpt=8)
    if foot: tip(D, s, foot, y=6.55)
    return s

def trainer(D, kicker, name, role, rows_, template=False):
    s = D.new(); header(D, s, kicker, "About the Trainer")
    card(s, 0.9, 1.95, 3.6, 3.9, fill=CHAR)
    circle(s, 2.25, 2.35, 0.9, ACC)
    text(s, 2.25, 2.35, 0.9, 0.9, "?" if template else "AA", 22, WHITE, bold=True, font=SERIF, align="c", anchor="m", autofit=False)
    text(s, 1.05, 3.45, 3.3, 0.6, name, 20, WHITE, bold=True, font=SERIF, align="c")
    text(s, 1.05, 4.1, 3.3, 1.2, role, 12.5, ACC, align="c")
    n = len(rows_); rh = 4.6 / n
    for i, (ic, lab, val) in enumerate(rows_):
        y = 1.95 + i * rh
        icon(s, ic, 5.0, y + 0.05, 0.5, fill=TEAL)
        if template:
            text(s, 5.7, y + 0.1, 7.0, 0.4, f"{lab}:  " + "_" * 48, 13, MUTED, autofit=False)
        else:
            rich(s, 5.7, y + 0.02, 7.0, rh - 0.05, [[(lab + "  ", dict(color=INK, bold=True, size=13)),
                                                     (val, dict(color=MUTED, size=13))]])
    return s

def flow(D, kicker, title, steps, sub=None, dark=False, tipt=None):
    """Five-step process flow (Assessment Flow / LMS download) — numbered cards with arrows."""
    s = D.new(dark=dark); header(D, s, kicker, title, sub)
    n = len(steps); gap = 0.3; cw = (12.1 - gap * (n - 1)) / n; top = 2.55 if sub else 2.2; ch = 2.9
    for i, (ic, t, d) in enumerate(steps):
        x = 0.6 + i * (cw + gap)
        card(s, x, top, cw, ch, fill=CHAR if (i % 2 == 0) == (not dark) and False else (PANEL if dark else WHITE))
        text(s, x + 0.22, top + 0.18, cw - 0.4, 0.3, f"STEP {i+1}", 11, ACC, bold=True, autofit=False)
        icon(s, ic, x + 0.22, top + 0.55, 0.56)
        text(s, x + 0.22, top + 1.25, cw - 0.4, 0.5, t, 15, WHITE if dark else INK, bold=True, font=SERIF, minpt=11)
        text(s, x + 0.22, top + 1.75, cw - 0.4, ch - 1.85, d, 12, SOFT if dark else MUTED, minpt=9)
        if i < n - 1:
            line(s, x + cw + 0.03, top + ch / 2, x + cw + gap - 0.03, top + ch / 2, ACC, 2, arrow=True)
    if tipt: tip(D, s, tipt, y=top + ch + 0.35)
    return s
